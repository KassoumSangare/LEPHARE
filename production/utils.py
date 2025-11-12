from .iautils import convert_to_date, unpack_ia_quotation_post_data
from openpyxl import load_workbook
import re
from collections import OrderedDict
from django.db import connection, transaction
import copy
from datetime import datetime
from decimal import Decimal
from customer.models import Client
from configuration_api.models import Qualite, Profession, TypeAssure
from .database import save_insured_ia, save_quotation_ia, enregistrer_ayant_droit
from .serializers import QuotationIaInsertionSerializer, DataInsertionSerializer

importation_col_list = [
    "Nom",
    "Prenoms",
    "NumeroCNI",
    "DateNaissance",
    "LieuNaissance",
    "Sexe",
    "NumeroTelephone",
    "NumeroMobile",
    "Qualite",
    "AdressePostale",
    "AdresseGeographique",
    "Email",
    "CapitalDeces",
    "CapitalInfirmite",
    "CapitalTraitement",
]


def check_date_format(date):
    regex = re.compile("[0-9]{4}\-[0-9]{2}\-[0-9]{2}")
    return re.match(regex, date)


def round_float_value(string):
    string = string.strip()
    if string.isdigit():
        return string
    elif string.replace(".", "", 1).isdigit():
        return str(round(float(string)))


def remove_unwanted_keys(data):
    for key in data.keys():
        if key not in importation_col_list:
            del data[key]

    return data


def load_people_data_from_excel(filename, index=0):
    dict_list = []
    # Les indices des lignes et des colonnes commencent par 1
    book = load_workbook(filename, data_only=True, keep_vba=False)
    sheet = book.worksheets[index]

    # Lecture de la ligne d'entête
    column_count = sheet.max_column

    keys = [
        str(sheet.cell(row=1, column=col_index).value).lower()
        for col_index in range(1, column_count + 1)
    ]

    # Read other lines from the second one
    row_count = sheet.max_row

    for row_index in range(2, row_count + 1):
        # empty = all(isinstance(cell, EmptyCell) for cell in row)
        d = OrderedDict(
            {
                keys[col_index - 1]: (
                    ""
                    if str(sheet.cell(row_index, col_index).value) == "None"
                    else sheet.cell(row_index, col_index).value
                )
                for col_index in range(1, column_count + 1)
            }
        )

        # Ignore blank lines
        if (str(d["nom"]).strip() in ("NA", "N/A", "N-A", "")) and (
            str(d["datenaissance"]).strip() in ("NA", "N/A", "N-A", "")
        ):
            continue

        # Remove unwanted data, especially blank colums
        # d = remove_unwanted_keys(d)

        # Keep only the date not the time
        # d["datenaissance"] = d["datenaissance"][:10]
        # print("Format de la date :", d["datenaissance"])

        dict_list.append(d)
    return dict_list


def insert_new_assure(data_line):
    output_code = -1
    output_msg = ""
    if data_line["qualite"] not in ("A", "a"):
        output_msg = "Cette personne n'est pas un assuré"
        return (output_code, output_msg)
    typeassure = TypeAssure.objects.get(pk=1)
    numerocompte = "NUMERO-COMPTE"
    creecie = "V"
    datenaissance = convert_to_date(data_line["datenaissance"])
    lieunaissance = str(data_line["lieunaissance"]).strip()
    statut = "V"
    numerocni = str(data_line["numerocni"]).strip()
    particulier = "V"
    email = str(data_line["email"]).strip()
    idprofession = 12
    sexe = str(data_line["sexe"]).strip().upper()
    if sexe == "M":
        idqualite = 1
    elif sexe == "F":
        idqualite = 2
    else:
        output_msg = "Sexe incorrect"
        return (output_code, output_msg)
    telephone = ""
    data_line["numerotelephone"] = str(data_line["numerotelephone"]).strip().upper()
    if data_line["numerotelephone"] not in ("NA", "N/A", "N-A"):
        telephone = data_line["numerotelephone"]
    mobile = ""
    data_line["numeromobile"] = str(data_line["numeromobile"]).strip().upper()
    if str(data_line["numeromobile"]) not in ("NA", "N/A", "N-A"):
        mobile = str(data_line["numeromobile"])

    adresse2 = ""
    data_line["adressegeographique"] = str(data_line["adressegeographique"]).strip()
    if str(data_line["adressegeographique"]).upper() not in ("NA", "N/A", "N-A"):
        adresse2 = str(data_line["adressegeographique"])

    adresse1 = ""
    data_line["adressepostale"] = str(data_line["adressepostale"]).strip()
    if str(data_line["adressepostale"]).upper() not in ("NA", "N/A", "N-A"):
        adresse1 = str(data_line["adressepostale"])
    vip = "V"
    prenoms = str(data_line["prenoms"]).strip()
    nom = str(data_line["nom"]).strip()
    client = Client(
        Nom=nom,
        Prenoms=prenoms,
        Vip=vip,
        Adresse1=adresse1,
        Adresse2=adresse2,
        Telephone=telephone,
        Mobile=mobile,
        IdQualite=idqualite,
        IdProfession=idprofession,
        Email=email,
        Particulier=particulier,
        CniPat=numerocni,
        Statut=statut,
        DateNaissance=datenaissance,
        LieuNaissance=lieunaissance,
        CreeCie=creecie,
        NumeroCompte=numerocompte,
        idtypeassure=typeassure,
    )
    client.save()
    output_code = client.IdClient
    output_msg = "Assuré enregistré avec succès"
    return (output_code, output_msg)


# Import insured people from Excel File
def import_ia_insured(filename, user_id, request_post_data):
    errors = []

    error_count = 0
    id_devis_initial = 0
    if "IdDevis" in request_post_data:
        if request_post_data["IdDevis"]:
            id_devis_initial = int(request_post_data["IdDevis"])
    try:
        request_data_list = load_people_data_from_excel(filename=filename, index=0)

        if request_data_list is None or len(request_data_list) == 0:
            error_count += 1
            errors.append("Le fichier Excel ne contient pas de données valides")
            return (error_count, id_devis, errors)

        insertion_data = copy.deepcopy(request_post_data)
        insertion_data["IdProduit"] = 2
        insertion_data["IdProfession"] = 0
        insertion_data["Flotte"] = (
            sum(
                1
                for d in request_data_list
                if str(d.get("qualite")).strip() in ("A", "a")
            )
        ) > 1
        insertion_data["Coassurance"] = False
        insertion_data["IdDevisDetail"] = 0
        insertion_data["CodeActivite"] = "01"

        # Parcourir les lignes extraites du fichier Excel en vue d'enregistrer les données
        id_assure = -1
        line_number = 1
        with transaction.atomic():
            for item in request_data_list:
                line_number += 1
                msg = ""
                if str(item["qualite"]).strip() in ("A", "a"):
                    (id_assure, msg) = insert_new_assure(item)
                    if id_assure and (id_assure != -1):
                        insertion_data["IdAssure"] = id_assure
                        insertion_data["CapitalDeces"] = item["capitaldeces"]
                        insertion_data["CapitalIpp"] = item["capitalinfirmite"]
                        insertion_data["FraisTraitement"] = item["capitaltraitement"]
                        insertion_data["DateNaissance"] = item["datenaissance"]
                        insertion_data["AdresseGeographique"] = item[
                            "adressegeographique"
                        ]
                        save_quotation_arg = unpack_ia_quotation_post_data(
                            insertion_data
                        )
                        with connection.cursor() as cursor:
                            cursor.execute(
                                "CALL sp_creation_devis_ia(%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);",
                                save_quotation_arg,
                            )
                            row = cursor.fetchone()
                            id_devis = row[0]
                            insertion_data["IdDevis"] = id_devis
                    else:
                        errors.append("Ligne n° " + str(line_number) + ":" + msg)
                        error_count += 1
                        transaction.rollback()
                        id_devis = id_devis_initial
                        break
                else:
                    if str(item["qualite"]).strip() in ("B", "b"):
                        lien = str(item["lien"]).strip()
                        if lien == "" or not lien.isdigit():
                            id_qualite_ayant_droit = 1
                        else:
                            id_qualite_ayant_droit = int(lien)
                        nom_ayant_droit = str(item["nom"]).strip()
                        prenoms_ayant_droit = str(item["prenoms"]).strip()
                        part_ayant_droit = Decimal(str(item["part"]).strip())
                        with connection.cursor() as cursor:
                            cursor.execute(
                                "CALL sp_saisie_ayant_droit_ia(%s, %s, %s, %s, %s, %s, %s);",
                                (
                                    id_assure,
                                    id_qualite_ayant_droit,
                                    nom_ayant_droit,
                                    prenoms_ayant_droit,
                                    part_ayant_droit,
                                    0,
                                    "",
                                ),
                            )
                    else:
                        errors.append(
                            "Ligne n° " + str(line_number) + ":" + "qualité incorrecte"
                        )
                        error_count += 1
                        transaction.rollback()
                        id_devis = id_devis_initial
                        break

    except Exception as error:
        print(error)
        error_count += 1
        errors.append("Ligne n° " + str(line_number) + ":" + str(error))
        transaction.rollback()
        id_devis = id_devis_initial

    return (error_count, id_devis, errors)
