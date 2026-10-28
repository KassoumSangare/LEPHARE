"""
Reprise des affaires antérieures : routes de l'écran Production > Affaires antérieures (production/reprise.py).

GET    /api/reprise/                      liste des reprises (filtres, totaux, pages de 50)
GET    /api/reprise/referentiels/         compagnies, branches et leurs catégories, énergies, marques
POST   /api/reprise/                      une affaire saisie à l'écran
DELETE /api/reprise/<id>/                 reprise erronée retirée (rien n'a été fait dessus)
GET    /api/reprise/modele/               modèle Excel de l'import
POST   /api/reprise/import/analyse/       fichier contrôlé ligne par ligne, rien n'est enregistré
POST   /api/reprise/import/               lignes correctes reprises, les autres signalées
"""

import io
import uuid
from datetime import date, datetime
from decimal import Decimal

from django.db import connection, transaction
from django.http import HttpResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from rest_framework import permissions, status
from rest_framework.authentication import BasicAuthentication
from rest_framework.decorators import api_view, authentication_classes, parser_classes, permission_classes
from rest_framework.parsers import FormParser, JSONParser, MultiPartParser
from rest_framework.response import Response

from autorisations.rbac import a_le_droit
from customer.models import Client
from institutionnel.authentication import KnoxOrDemoTokenAuthentication

from .reprise import (PAIEMENTS, ErreurReprise, controler, description_devis, devis_a_finaliser, devis_reprise, enregistrer,
                      marquer_devis, normaliser, reprendre, reprendre_devis, risque_branche, supprimer)
from .views import ContratViewSet

AUTHENTIFICATION = [KnoxOrDemoTokenAuthentication, BasicAuthentication]
PAGE = 50
TAILLE_MAX = 2000
LIGNES_MAX = 2000

# Colonnes du modèle Excel : (titre, clé de l'API ou de la résolution, aide)
COLONNES = [
    ("N° client", "IdClient", "Facultatif : numéro du client dans LE PHARE. Sinon le client est cherché par son nom (et son téléphone), créé s'il n'existe pas."),
    ("Nom ou raison sociale", "Nom", "Obligatoire sans n° client."),
    ("Prénoms", "Prenoms", ""),
    ("Téléphone", "Telephone", "Départage les homonymes ; repris sur la fiche d'un nouveau client."),
    ("Type de client", "TypeClient", "Particulier (par défaut) ou Entreprise."),
    ("Compagnie", "Compagnie", "Raison sociale, ex. NSIA ASSURANCES (ou son début)."),
    ("Branche", "Branche", "Automobile, Individuelle Accidents, Voyage, MRH, Santé, Transport, MRP, RC, Dommages."),
    ("Catégorie", "Categorie", "Facultatif : code de la catégorie (ex. 201) ; la première de la branche à défaut."),
    ("N° police compagnie", "NumeroPoliceCompagnie", "Obligatoire."),
    ("Date d'émission", "DateEmission", "jj/mm/aaaa, avant le 1er janvier de l'année en cours."),
    ("Date d'effet", "DateEffet", "jj/mm/aaaa, avant le 1er janvier de l'année en cours."),
    ("Date d'échéance", "DateExpiration", "jj/mm/aaaa."),
    ("Prime nette", "PrimeNette", "Automobile : FGA et CEDEAO compris, comme sur la quittance."),
    ("Accessoires", "Accessoire", ""),
    ("Taxes", "Taxe", ""),
    ("FGA", "Fga", "Automobile, facultatif (compris dans la prime nette)."),
    ("CEDEAO", "Cedeao", "Automobile, facultatif (comprise dans la prime nette)."),
    ("Prime TTC", "PrimeTtc", "Prime nette + accessoires + taxes ; calculée si vide."),
    ("Taux de commission (%)", "TauxCommission", "Facultatif : taux de la compagnie et de la branche à défaut. "
                                                    "La quittance est reprise encaissée et la commission perçue."),
    ("Immatriculation", "Immatriculation", "Automobile : obligatoire."),
    ("Marque", "Marque", "Automobile."),
    ("Modèle", "Modele", "Automobile."),
    ("Puissance (CV)", "Puissance", "Automobile."),
    ("Places", "Places", "Automobile."),
    ("Énergie", "Energie", "Essence, Diésel, Électrique, Hybride."),
    ("Valeur neuve", "ValeurNeuve", ""),
    ("Valeur vénale", "ValeurVenale", ""),
    ("Date de mise en circulation", "DateMec", "jj/mm/aaaa."),
    ("Description du risque", "Risque", "Autres branches : objet assuré, lieu, capitaux…"),
    ("Observation", "Observation", ""),
]
SYNONYMES_BRANCHES = {
    "AUTO": 1, "AUTOMOBILE": 1, "IA": 2, "INDIVIDUELLE": 2, "INDIVIDUELLE ACCIDENT": 2, "INDIVIDUELLE ACCIDENTS": 2,
    "VOYAGE": 3, "ASSURANCE VOYAGE": 3, "MRH": 4, "MULTIRISQUE HABITATION": 4, "MULTIRISQUES HABITATION": 4, "SANTE": 5,
    "TRANSPORT": 6, "MRP": 7, "MULTIRISQUE PROFESSIONNELLE": 7, "MULTIRISQUES PROFESSIONNELLE": 7, "RC": 8,
    "RESPONSABILITE CIVILE": 8, "DOMMAGES": 9, "TOUS DOMMAGES": 9,
}


def _lignes_sql(cursor, sql, params=None):
    cursor.execute(sql, params or [])
    colonnes = [c[0] for c in cursor.description]
    return [dict(zip(colonnes, ligne)) for ligne in cursor.fetchall()]


def _iso(v):
    return v.date().isoformat() if isinstance(v, datetime) else v.isoformat() if isinstance(v, date) else v


def _erreur(message, code=status.HTTP_400_BAD_REQUEST):
    return Response({"error": message}, status=code)


# Affaire renouvelée : renouvellement (avenant REN) confirmé en contrat, non annulé, sur la même police. Elle est alors
# migrée dans le module Contrats et quitte la liste des affaires antérieures (demande du 2026-10-27 : « après
# renouvellement ils doivent être migrés dans le module contrat ») ; ContratViewSet._filtres_portefeuille fait l'inverse.
SQL_RENOUVELEE = """EXISTS (SELECT 1 FROM stdcontrat x JOIN stdavenant xa ON xa.idavenant = x.idavenant
                            WHERE x.numeropolice = c.numeropolice AND xa.codeavenant = 'REN'
                              AND COALESCE(x.idcontratannulation, 0) = 0 AND x.dateeffet > c.dateeffet)"""


# ------------------------------------------------------------------ liste
@api_view(["GET", "POST"])
@authentication_classes(AUTHENTIFICATION)
@permission_classes([permissions.IsAuthenticated])
def reprises(request):
    """GET : liste filtrée (annee, compagnie, produit, q, page) des affaires pas encore renouvelées, nombre d'affaires
    renouvelées (migrées dans le module Contrats) ; POST : une affaire."""
    if request.method == "POST":
        try:
            resultat = reprendre(request.data, request.user)
        except ErreurReprise as erreur:
            return _erreur(str(erreur))
        return Response({**resultat, "message": f"Affaire reprise : police {resultat['NumeroPolice']}, quittance "
                                                 f"{resultat['NumeroQuittance']}."}, status=status.HTTP_201_CREATED)
    p = request.query_params
    conditions, params = ["TRUE"], []
    if p.get("annee"):
        conditions.append("EXTRACT(YEAR FROM c.dateeffet) = %s")
        params.append(int(p["annee"]))
    if p.get("compagnie"):
        conditions.append("c.idcompagnie = %s")
        params.append(int(p["compagnie"]))
    if p.get("q"):
        conditions.append("(UPPER(c.numeropolice) LIKE %s OR UPPER(COALESCE(c.numeropolicecompagnie, '')) LIKE %s "
                          "OR UPPER(TRIM(cl.nom || ' ' || COALESCE(cl.prenoms, ''))) LIKE %s OR UPPER(q.numeroquittance) LIKE %s)")
        terme = f"%{str(p['q']).strip().upper()}%"
        params += [terme] * 4
    # Branche : un produit, ou plusieurs séparés par des virgules (onglet du registre regroupant des produits)
    produits = [int(x) for x in str(p.get("produit") or "").split(",") if x.strip().isdigit()]
    # Affaires renouvelées, migrées dans le module Contrats : comptées à part, hors de la liste
    conditions_migrees = conditions + [SQL_RENOUVELEE] + (["c.idproduit = ANY(%s)"] if produits else [])
    params_migrees = params + ([produits] if produits else [])
    conditions.append(f"NOT {SQL_RENOUVELEE}")
    # Compteurs par branche (onglets) : mêmes filtres, branche exceptée
    conditions_branches, params_branches = list(conditions), list(params)
    if produits:
        conditions.append("c.idproduit = ANY(%s)")
        params.append(produits)

    def base(where):
        return f"""FROM lephare_reprise r
                   JOIN stdcontrat c ON c.idcontrat = r.idcontrat
                   JOIN stdquittance q ON q.idquittance = r.idquittance
                   LEFT JOIN stdclient cl ON cl.idclient = c.idclient
                   LEFT JOIN stdcompagnie co ON co.idcompagnie = c.idcompagnie
                   LEFT JOIN stdproduit pr ON pr.idproduit = c.idproduit
                   WHERE {' AND '.join(where)}"""
    page = max(1, int(p.get("page") or 1))
    # Taille de page : 50 par défaut ; la vue (tableau paginé côté écran, comme la Gestion des devis) en demande 2 000
    taille = min(max(int(p.get("taille") or PAGE), 1), TAILLE_MAX)
    with connection.cursor() as cursor:
        # Affaires antérieures soldées (consigne du 2026-10-27) : primes encaissées, commissions perçues
        totaux = _lignes_sql(cursor, f"""SELECT COUNT(*) AS nombre, COALESCE(SUM(q.primettc), 0) AS ttc,
                                               COALESCE(SUM(c.primenette), 0) AS primenette,
                                               COALESCE(SUM(q.mt_encaisse), 0) AS encaisse,
                                               COALESCE(SUM(c.commissionintermediaire), 0) AS commission
                                        {base(conditions)}""", params)[0]
        migrees = _lignes_sql_brut(cursor, f"SELECT COUNT(*) {base(conditions_migrees)}", params_migrees)[0][0]
        lignes = _lignes_sql(cursor, f"""SELECT r.id, r.idcontrat, r.mode, r.lot, r.commission_percue, r.creepar, r.idutilisateur,
                                                r.datecreation, c.numeropolice, c.numeropolicecompagnie, c.dateemission, c.dateeffet,
                                                c.dateexpiration, c.idcompagnie, co.raisonsociale AS compagnie, c.idproduit,
                                                pr.libelleproduit AS branche, c.idclient,
                                                TRIM(cl.nom || ' ' || COALESCE(cl.prenoms, '')) AS client,
                                                c.commissionintermediaire AS commission, c.taux_commission,
                                                q.numeroquittance, q.primettc, COALESCE(q.mt_encaisse, 0) AS encaisse,
                                                r.donnees ->> 'Description' AS description,
                                                -- Échéance, comme les statuts du portefeuille des contrats
                                                CASE WHEN COALESCE(c.idcontratannulation, 0) <> 0 THEN 'annule'
                                                     WHEN c.dateexpiration::date < CURRENT_DATE THEN 'echu'
                                                     WHEN c.dateexpiration::date <= CURRENT_DATE + %s THEN 'a_renouveler'
                                                     ELSE 'en_cours' END AS statut,
                                                -- Devis de renouvellement ouvert, pas encore confirmé (production/renouvellement.py)
                                                (SELECT d.numerodevis FROM stddevis d JOIN stdavenant da ON da.idavenant = d.idavenant
                                                 WHERE da.codeavenant = 'REN' AND d.idoldhist = c.iddevis
                                                   AND NOT COALESCE(d.confirme, false) AND NOT COALESCE(d.archive, false)
                                                 ORDER BY d.iddevis DESC LIMIT 1) AS devis_renouvellement
                                         {base(conditions)}
                                         ORDER BY r.id DESC LIMIT {taille} OFFSET %s""",
                             [ContratViewSet.DELAI_RENOUVELLEMENT_JOURS] + params + [(page - 1) * taille])
        par_branche = {str(i): n for i, n in _lignes_sql_brut(cursor, f"""
            SELECT c.idproduit, COUNT(*) FROM lephare_reprise r
            JOIN stdcontrat c ON c.idcontrat = r.idcontrat JOIN stdquittance q ON q.idquittance = r.idquittance
            LEFT JOIN stdclient cl ON cl.idclient = c.idclient
            WHERE {' AND '.join(conditions_branches)} GROUP BY 1""", params_branches)}
        annees = [a[0] for a in _lignes_sql_brut(cursor, "SELECT DISTINCT EXTRACT(YEAR FROM c.dateeffet)::int FROM lephare_reprise r "
                                                         f"JOIN stdcontrat c ON c.idcontrat = r.idcontrat WHERE NOT {SQL_RENOUVELEE} "
                                                         "ORDER BY 1 DESC")]
    resultat = []
    for l in lignes:
        ttc, encaisse = Decimal(l["primettc"] or 0), Decimal(l["encaisse"] or 0)
        statut = "PAYEE" if encaisse >= ttc else "PARTIELLE" if encaisse > 0 else "IMPAYEE"
        resultat.append({
            "Id": l["id"], "IdContrat": l["idcontrat"], "NumeroPolice": l["numeropolice"], "NumeroPoliceCompagnie": l["numeropolicecompagnie"] or "",
            "NumeroQuittance": l["numeroquittance"], "Client": l["client"] or "", "IdClient": l["idclient"],
            "Compagnie": l["compagnie"] or "", "Branche": l["branche"] or "", "DateEmission": _iso(l["dateemission"]),
            "DateEffet": _iso(l["dateeffet"]), "DateExpiration": _iso(l["dateexpiration"]), "PrimeTtc": float(ttc),
            "Encaisse": float(encaisse), "Reste": float(ttc - encaisse), "Paiement": statut, "PaiementLibelle": PAIEMENTS[statut],
            "Commission": float(l["commission"] or 0), "TauxCommission": float(l["taux_commission"] or 0),
            "CommissionPercue": l["commission_percue"], "Mode": l["mode"], "Lot": l["lot"] or "", "CreePar": l["creepar"],
            "Auteur": l["idutilisateur"] == request.user.id, "DateCreation": _iso(l["datecreation"]),
            "IdProduit": l["idproduit"], "Description": l["description"] or "",
            "Statut": l["statut"], "DevisRenouvellement": l["devis_renouvellement"] or "",
        })
    return Response({
        "Lignes": resultat, "Page": page, "Pages": max(1, -(-int(totaux["nombre"]) // taille)), "Annees": annees,
        "ParBranche": par_branche, "Migrees": int(migrees),
        "Totaux": {"Nombre": int(totaux["nombre"]), "PrimeTtc": float(totaux["ttc"]), "PrimeNette": float(totaux["primenette"]),
                   "Encaisse": float(totaux["encaisse"]), "Commission": float(totaux["commission"])},
        "PeutToutSupprimer": a_le_droit(request.user, "contrats.cancel"),
    })


def _lignes_sql_brut(cursor, sql, params=None):
    cursor.execute(sql, params or [])
    return cursor.fetchall()


SQL_REPRISE = """
    SELECT r.id, r.idcontrat, r.mode, r.lot, r.donnees, r.commission_percue, r.date_commission_percue, r.creepar,
           r.datecreation, c.numeropolice, c.numeropolicecompagnie, c.idproduit, pr.libelleproduit AS branche,
           co.raisonsociale AS compagnie, TRIM(cl.nom || ' ' || COALESCE(cl.prenoms, '')) AS client, c.idclient,
           c.dateemission, c.dateeffet, c.dateexpiration, c.primenette, c.accessoire, c.taxe, c.fga, c.cedeao, c.primettc,
           c.taux_commission, c.commissionintermediaire, q.numeroquittance, COALESCE(q.mt_encaisse, 0) AS encaisse,
           t.libelle AS categorie
    FROM lephare_reprise r
    JOIN stdcontrat c ON c.idcontrat = r.idcontrat
    JOIN stdquittance q ON q.idquittance = r.idquittance
    LEFT JOIN stdclient cl ON cl.idclient = c.idclient
    LEFT JOIN stdcompagnie co ON co.idcompagnie = c.idcompagnie
    LEFT JOIN stdproduit pr ON pr.idproduit = c.idproduit
    LEFT JOIN stdcontratdetail cd ON cd.idcontrat = c.idcontrat
    LEFT JOIN stdtarif t ON t.idtarif = cd.idtarif
"""


def _fiche(ligne):
    """Reprise lue en base : contrat, primes et risque tels que saisis, libellés des champs de la branche."""
    donnees = ligne["donnees"] if isinstance(ligne["donnees"], dict) else {}
    risque = risque_branche(ligne["idproduit"])
    ttc, encaisse = Decimal(ligne["primettc"] or 0), Decimal(ligne["encaisse"] or 0)
    statut = "PAYEE" if encaisse >= ttc else "PARTIELLE" if encaisse > 0 else "IMPAYEE"
    return {
        "Id": ligne["id"], "IdContrat": ligne["idcontrat"], "NumeroPolice": ligne["numeropolice"],
        "NumeroPoliceCompagnie": ligne["numeropolicecompagnie"] or "", "NumeroQuittance": ligne["numeroquittance"],
        "IdProduit": ligne["idproduit"], "Branche": ligne["branche"] or "", "Categorie": (ligne["categorie"] or "").strip(),
        "Compagnie": ligne["compagnie"] or "", "Client": ligne["client"] or "", "IdClient": ligne["idclient"],
        "DateEmission": _iso(ligne["dateemission"]), "DateEffet": _iso(ligne["dateeffet"]), "DateExpiration": _iso(ligne["dateexpiration"]),
        "PrimeNette": float(ligne["primenette"] or 0), "Accessoire": float(ligne["accessoire"] or 0), "Taxe": float(ligne["taxe"] or 0),
        "Fga": float(ligne["fga"] or 0), "Cedeao": float(ligne["cedeao"] or 0), "PrimeTtc": float(ttc),
        "TauxCommission": float(ligne["taux_commission"] or 0), "Commission": float(ligne["commissionintermediaire"] or 0),
        "CommissionPercue": ligne["commission_percue"], "DateCommissionPercue": _iso(ligne["date_commission_percue"]),
        "Encaisse": float(encaisse), "Reste": float(ttc - encaisse), "Paiement": statut, "PaiementLibelle": PAIEMENTS[statut],
        "Etape": risque["Etape"], "Description": donnees.get("Description") or "",
        "Champs": [{"Libelle": l, "Valeur": val} for l, val in donnees.get("Affichage") or []],
        "Precisions": donnees.get("Risque") or "", "Observation": donnees.get("Observation") or "",
        "Mode": ligne["mode"], "Lot": ligne["lot"] or "", "CreePar": ligne["creepar"], "DateCreation": _iso(ligne["datecreation"]),
    }


def _lire(condition, valeur):
    with connection.cursor() as cursor:
        lignes = _lignes_sql(cursor, SQL_REPRISE + f" WHERE {condition} = %s ORDER BY cd.idcontratdetail LIMIT 1", [valeur])
    if not lignes:
        return None
    ligne = lignes[0]
    if isinstance(ligne["donnees"], str):
        import json
        ligne["donnees"] = json.loads(ligne["donnees"])
    return _fiche(ligne)


@api_view(["GET"])
@authentication_classes(AUTHENTIFICATION)
@permission_classes([permissions.IsAuthenticated])
def reprise_du_contrat(request, id_contrat):
    """GET : reprise d'un contrat (fiche contrat), 404 si le contrat n'a pas été repris."""
    fiche = _lire("r.idcontrat", id_contrat)
    return Response(fiche) if fiche else _erreur("Ce contrat n'est pas une affaire antérieure reprise.", status.HTTP_404_NOT_FOUND)


@api_view(["GET", "DELETE"])
@authentication_classes(AUTHENTIFICATION)
@permission_classes([permissions.IsAuthenticated])
def reprise(request, pk):
    """GET : fiche de la reprise ; DELETE : reprise erronée retirée."""
    if request.method == "GET":
        fiche = _lire("r.id", pk)
        return Response(fiche) if fiche else _erreur("Reprise introuvable.", status.HTTP_404_NOT_FOUND)
    try:
        resultat = supprimer(pk, request.user, a_le_droit(request.user, "contrats.cancel"))
    except ErreurReprise as erreur:
        return _erreur(str(erreur))
    return Response({**resultat, "message": f"Reprise retirée : police {resultat['NumeroPolice']} et quittance {resultat['NumeroQuittance']} supprimées."})


def _devis_json(cursor, v):
    """Devis saisi pour une reprise : identité, primes et risque (finalisation, liste des devis à finaliser)."""
    return {
        "IdDevis": v["iddevis"], "NumeroDevis": v["numerodevis"], "IdProduit": v["idproduit"], "Branche": v["branche"] or "",
        "IdCompagnie": v["idcompagnie"], "Compagnie": v["compagnie"] or "", "IdClient": v["idclient"], "Client": v["client"] or "",
        "DateEmission": _iso(v["dateemission"]), "DateEffet": _iso(v["dateeffet"]), "DateExpiration": _iso(v["dateexpiration"]),
        "PrimeNette": float(v["primenette"] or 0), "Accessoire": float(v["accessoire"] or 0), "Taxe": float(v["taxe"] or 0),
        "Fga": float(v["fga"] or 0), "Cedeao": float(v["cedeao"] or 0), "PrimeTtc": float(v["primettc"] or 0),
        "NumeroPoliceCompagnie": v["numeropolicecompagnie"] or "", "Flotte": v["flotte"], "DateCreation": _iso(v["datecreation"]),
        "Champs": [{"Libelle": l, "Valeur": val} for l, val in description_devis(cursor, v["iddevis"], v["idproduit"])],
    }


@api_view(["GET"])
@authentication_classes(AUTHENTIFICATION)
@permission_classes([permissions.IsAuthenticated])
def devis_en_attente(request):
    """GET : devis saisis pour une reprise (effet avant le 1er janvier, enregistrés cette année), pas encore finalisés."""
    with connection.cursor() as cursor:
        return Response({"Devis": [_devis_json(cursor, v) for v in devis_a_finaliser(cursor)]})


@api_view(["POST"])
@authentication_classes(AUTHENTIFICATION)
@permission_classes([permissions.IsAuthenticated])
def devis_attente(request, id_devis):
    """POST : devis enregistré en mode reprise contrôlé (affaire nouvelle, effet avant le 1er janvier) et noté en attente
    de finalisation ; 400 avec le motif sinon (fenêtre du devis enregistré)."""
    try:
        v = marquer_devis(id_devis, request.user)
    except ErreurReprise as erreur:
        return _erreur(str(erreur))
    return Response({"IdDevis": v["iddevis"], "NumeroDevis": v["numerodevis"]})


@api_view(["GET", "POST"])
@authentication_classes(AUTHENTIFICATION)
@permission_classes([permissions.IsAuthenticated])
def devis(request, id_devis):
    """GET : devis à finaliser ; POST : devis confirmé en affaire antérieure (police compagnie, date d'émission d'origine),
    quittance encaissée et commission perçue."""
    if request.method == "GET":
        try:
            with connection.cursor() as cursor:
                return Response(_devis_json(cursor, devis_reprise(cursor, id_devis)))
        except ErreurReprise as erreur:
            return _erreur(str(erreur))
    try:
        resultat = reprendre_devis(id_devis, request.data, request.user)
    except ErreurReprise as erreur:
        return _erreur(str(erreur))
    return Response({**resultat, "message": f"Affaire antérieure enregistrée : police {resultat['NumeroPolice']}, quittance "
                                             f"{resultat['NumeroQuittance']}."}, status=status.HTTP_201_CREATED)


@api_view(["GET"])
@authentication_classes(AUTHENTIFICATION)
@permission_classes([permissions.IsAuthenticated])
def referentiels(request):
    """Listes de l'écran : compagnies actives, branches actives (catégories, étape du risque et ses champs), énergies,
    marques et genres de véhicule."""
    with connection.cursor() as cursor:
        compagnies = _lignes_sql(cursor, "SELECT idcompagnie AS id, raisonsociale AS nom FROM stdcompagnie "
                                         "WHERE idcompagnie > 0 AND COALESCE(active, true) ORDER BY raisonsociale")
        produits = _lignes_sql(cursor, "SELECT idproduit AS id, libelleproduit AS libelle FROM stdproduit "
                                       "WHERE active AND idproduit BETWEEN 1 AND 9 ORDER BY idproduit")
        for p in produits:
            p.update(risque_branche(p["id"]))
            p["categories"] = [{"id": t, "libelle": f"{c} - {l}", "code": c}
                               for t, l, c in _lignes_sql_brut(cursor, "SELECT id_tarif, libelle_tarif, TRIM(code_categorie) "
                                                                       "FROM fn_liste_tarif_produit(%s) ORDER BY 3, 1", [p["id"]])]
        energies = _lignes_sql(cursor, "SELECT idenergie AS id, libelle FROM stdenergie WHERE idenergie > 0 ORDER BY 1")
        marques = _lignes_sql(cursor, "SELECT idmarque AS id, libellemarque AS libelle FROM stdmarque ORDER BY libellemarque")
        genres = _lignes_sql(cursor, "SELECT idgenre AS id, libellegenre AS libelle FROM stdgenrevehicule WHERE idgenre > 0 ORDER BY 1")
    return Response({"Compagnies": compagnies, "Produits": produits, "Energies": energies, "Marques": marques, "Genres": genres,
                     "AnneeLimite": date.today().year})


# ------------------------------------------------------------------ import Excel
@api_view(["GET"])
@authentication_classes(AUTHENTIFICATION)
@permission_classes([permissions.IsAuthenticated])
def modele(request):
    """Modèle Excel : une ligne par quittance reprise, aide de chaque colonne en deuxième feuille."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Affaires antérieures"
    ws.append([t for t, _, _ in COLONNES])
    ws.append(["", "KOUASSI", "AMANI JEAN", "0700000000", "Particulier", "NSIA ASSURANCES", "Automobile", "201", "PC-2023-001245",
               "15/03/2023", "15/03/2023", "14/03/2024", 98205, 5000, 14965, "", 1000, 118170, 20,
               "AB 1234 CI", "TOYOTA", "COROLLA", 8, 5, "Essence", 15000000, 9000000, "10/01/2019", "", "Ligne d'exemple : à remplacer"])
    for cellule in ws[1]:
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = PatternFill("solid", fgColor="1F4E78")
        cellule.alignment = Alignment(wrap_text=True, vertical="center")
    for i, (titre, _, _) in enumerate(COLONNES, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(14, min(30, len(titre) + 4))
    ws.freeze_panes = "A2"
    aide = wb.create_sheet("Aide")
    aide.append(["Colonne", "Contenu attendu"])
    for titre, _, texte in COLONNES:
        aide.append([titre, texte])
    aide.column_dimensions["A"].width = 30
    aide.column_dimensions["B"].width = 110
    for cellule in aide[1]:
        cellule.font = Font(bold=True)
    tampon = io.BytesIO()
    wb.save(tampon)
    reponse = HttpResponse(tampon.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    reponse["Content-Disposition"] = 'attachment; filename="modele_reprise_affaires_anterieures.xlsx"'
    return reponse


def _lire_fichier(fichier):
    """Lignes du fichier (dict par clé de COLONNES, numéro de ligne Excel). Lève ErreurReprise."""
    try:
        wb = load_workbook(fichier, data_only=True, read_only=True)
    except Exception as erreur:  # fichier illisible
        raise ErreurReprise(f"Fichier Excel illisible : {str(erreur).splitlines()[0]}") from erreur
    ws = wb.worksheets[0]
    lignes = list(ws.iter_rows(values_only=True))
    if not lignes:
        raise ErreurReprise("Le fichier est vide.")
    entetes = [normaliser(str(v or "")).replace("'", " ") for v in lignes[0]]
    index = {}
    for titre, cle, _ in COLONNES:
        t = normaliser(titre).replace("'", " ")
        if t in entetes:
            index[cle] = entetes.index(t)
    manquantes = [titre for titre, cle, _ in COLONNES if cle in ("Compagnie", "Branche", "NumeroPoliceCompagnie", "DateEffet",
                                                                  "DateExpiration", "PrimeNette") and cle not in index]
    if manquantes:
        raise ErreurReprise("Colonnes absentes de la 1re ligne : " + ", ".join(manquantes) + " (téléchargez le modèle).")
    resultat = []
    for numero, ligne in enumerate(lignes[1:], start=2):
        valeurs = {cle: ligne[i] if i < len(ligne) else None for cle, i in index.items()}
        if all(v in (None, "") for v in valeurs.values()):
            continue
        if "ligne d'exemple" in str(valeurs.get("Observation") or "").lower():
            continue
        valeurs["_ligne"] = numero
        resultat.append(valeurs)
    if len(resultat) > LIGNES_MAX:
        raise ErreurReprise(f"{LIGNES_MAX} lignes au plus par fichier : découpez-le.")
    return resultat


class _Resolution:
    """Compagnies, branches, catégories, énergies et marques désignées en clair dans le fichier ; clients retrouvés."""

    def __init__(self, cursor):
        self.cursor = cursor
        self.compagnies = [(i, normaliser(n)) for i, n in _lignes_sql_brut(cursor, "SELECT idcompagnie, raisonsociale FROM stdcompagnie WHERE idcompagnie > 0")]
        self.produits = {i: normaliser(l) for i, l in _lignes_sql_brut(cursor, "SELECT idproduit, libelleproduit FROM stdproduit WHERE idproduit BETWEEN 1 AND 9")}
        self.energies = {normaliser(l): i for i, l in _lignes_sql_brut(cursor, "SELECT idenergie, libelle FROM stdenergie")}
        self.energies.update({"DIESEL": 2, "GASOIL": 2, "GAZOIL": 2})
        self.marques = {normaliser(l): i for i, l in _lignes_sql_brut(cursor, "SELECT idmarque, libellemarque FROM stdmarque")}
        self.nouveaux = {}

    def compagnie(self, valeur, id_produit=None):
        """Raison sociale exacte, sinon son début (« NSIA ») ; une affaire non-vie écarte les compagnies vie
        (NSIA VIE, ALLIANZ CI VIE) quand le nom désigne aussi la compagnie IARD."""
        texte = normaliser(valeur)
        if not texte:
            raise ErreurReprise("Compagnie obligatoire.")
        if texte.isdigit() and any(i == int(texte) for i, _ in self.compagnies):
            return int(texte)
        exactes = [i for i, n in self.compagnies if n == texte]
        trouvees = exactes or [i for i, n in self.compagnies if n.startswith(texte) or texte in n]
        if len(trouvees) > 1 and id_produit and id_produit <= 9:
            non_vie = [i for i in trouvees if " VIE" not in f" {dict(self.compagnies)[i]} "]
            trouvees = non_vie or trouvees
        if len(trouvees) != 1:
            noms = ", ".join(dict(self.compagnies)[i] for i in trouvees[:4])
            raise ErreurReprise(f"Compagnie « {valeur} » inconnue." if not trouvees else f"Compagnie « {valeur} » ambiguë ({noms}) : précisez.")
        return trouvees[0]

    def produit(self, valeur):
        texte = normaliser(valeur)
        if texte in SYNONYMES_BRANCHES:
            return SYNONYMES_BRANCHES[texte]
        trouves = [i for i, l in self.produits.items() if l == texte or l.startswith(texte)]
        if len(trouves) != 1:
            raise ErreurReprise(f"Branche « {valeur} » inconnue.")
        return trouves[0]

    def categorie(self, id_produit, valeur):
        texte = normaliser(valeur)
        if not texte:
            return None
        for t, l, c in _lignes_sql_brut(self.cursor, "SELECT id_tarif, libelle_tarif, TRIM(code_categorie) FROM fn_liste_tarif_produit(%s) "
                                                      "ORDER BY 3, 1", [id_produit]):
            if texte in (c, normaliser(l)):
                return t
        raise ErreurReprise(f"Catégorie « {valeur} » inconnue pour cette branche.")

    def client(self, v):
        if str(v.get("IdClient") or "").strip():
            return int(float(str(v["IdClient"]).strip()))
        nom = normaliser(v.get("Nom"))
        if not nom:
            raise ErreurReprise("Client : n° client ou nom obligatoire.")
        prenoms = normaliser(v.get("Prenoms"))
        telephone = "".join(ch for ch in str(v.get("Telephone") or "") if ch.isdigit())
        cle = (nom, prenoms, telephone)
        if cle in self.nouveaux:
            return self.nouveaux[cle]
        candidats = [(i, t, m) for i, n, p, t, m in _lignes_sql_brut(
            self.cursor, "SELECT idclient, nom, prenoms, telephone, mobile FROM stdclient WHERE statut <> 'A' "
                         "AND UPPER(TRIM(nom)) = %s", [str(v.get("Nom")).strip().upper()]) if normaliser(n) == nom and normaliser(p) == prenoms]
        if telephone and len(candidats) > 1:
            candidats = [c for c in candidats if telephone in "".join(ch for ch in f"{c[1] or ''} {c[2] or ''}" if ch.isdigit())] or candidats
        if len(candidats) > 1:
            raise ErreurReprise(f"{len(candidats)} clients « {nom} {prenoms} » : indiquez le n° client "
                                f"({', '.join(str(c[0]) for c in candidats[:5])}).")
        if candidats:
            return candidats[0][0]
        entreprise = normaliser(v.get("TypeClient")).startswith("ENTREPRISE")
        client = Client.objects.create(
            Nom=str(v.get("Nom")).strip()[:150], Prenoms=str(v.get("Prenoms") or "").strip()[:150] or None,
            Telephone=str(v.get("Telephone") or "").strip()[:20], Mobile=str(v.get("Telephone") or "").strip()[:20],
            Particulier="F" if entreprise else "V", Statut="V", Vip="V", CreeCie="V", NumeroCompte="NUMERO-COMPTE",
            Adresse1="", Adresse2="", Email="", Fonction="",
        )
        self.nouveaux[cle] = client.IdClient
        return client.IdClient

    def donnees(self, v):
        d = {k: v.get(k) for _, k, _ in COLONNES}
        d["IdProduit"] = self.produit(v.get("Branche"))
        d["IdCompagnie"] = self.compagnie(v.get("Compagnie"), d["IdProduit"])
        d["IdClient"] = self.client(v)
        d["IdTarif"] = self.categorie(d["IdProduit"], v.get("Categorie"))
        if str(v.get("Energie") or "").strip():
            d["IdEnergie"] = self.energies.get(normaliser(v.get("Energie")))
            if not d["IdEnergie"]:
                raise ErreurReprise(f"Énergie « {v.get('Energie')} » inconnue.")
        if str(v.get("Marque") or "").strip():
            d["IdMarque"] = self.marques.get(normaliser(v.get("Marque"))) or 0
            if not d["IdMarque"]:
                d["Modele"] = f"{v.get('Marque')} {v.get('Modele') or ''}".strip()
        return d


def _importer(fichier, utilisateur, executer):
    """Toutes les lignes reprises dans une transaction (point de sauvegarde par ligne) : les lignes refusées sont
    signalées ; l'analyse annule tout à la fin, l'import garde les lignes correctes."""
    lignes = _lire_fichier(fichier)
    lot = uuid.uuid4().hex[:12]
    compte_rendu = []
    with transaction.atomic(), connection.cursor() as cursor:
        resolution = _Resolution(cursor)
        for v in lignes:
            point = transaction.savepoint()
            nouveaux_avant = len(resolution.nouveaux)
            try:
                donnees = resolution.donnees(v)
                r = controler(cursor, donnees)
                cree = enregistrer(cursor, r, utilisateur, mode="import", lot=lot)
                transaction.savepoint_commit(point)
                compte_rendu.append({
                    "Ligne": v["_ligne"], "Statut": "OK", "Message": "",
                    "Client": f"{v.get('Nom') or ''} {v.get('Prenoms') or ''}".strip() or f"n° {donnees['IdClient']}",
                    "NouveauClient": len(resolution.nouveaux) > nouveaux_avant, "NumeroPoliceCompagnie": r["NumeroPoliceCompagnie"],
                    "DateEffet": r["DateEffet"].isoformat(), "PrimeTtc": float(r["PrimeTtc"]), "Paiement": PAIEMENTS[r["StatutPaiement"]],
                    "NumeroPolice": cree["NumeroPolice"], "NumeroQuittance": cree["NumeroQuittance"],
                })
            except (ErreurReprise, ValueError) as erreur:
                transaction.savepoint_rollback(point)
                resolution.nouveaux = dict(list(resolution.nouveaux.items())[:nouveaux_avant])
                compte_rendu.append({"Ligne": v["_ligne"], "Statut": "ERREUR", "Message": str(erreur),
                                     "Client": f"{v.get('Nom') or ''} {v.get('Prenoms') or ''}".strip(),
                                     "NumeroPoliceCompagnie": str(v.get("NumeroPoliceCompagnie") or ""), "PrimeTtc": None})
        if not executer:
            transaction.set_rollback(True)
    ok = [l for l in compte_rendu if l["Statut"] == "OK"]
    return {
        "Lot": lot if executer else "", "Lignes": compte_rendu, "Correctes": len(ok), "Erreurs": len(compte_rendu) - len(ok),
        "NouveauxClients": sum(1 for l in ok if l.get("NouveauClient")), "PrimeTtc": sum(l["PrimeTtc"] for l in ok),
        "Importe": executer,
    }


@api_view(["POST"])
@authentication_classes(AUTHENTIFICATION)
@permission_classes([permissions.IsAuthenticated])
@parser_classes([MultiPartParser, FormParser, JSONParser])
def import_analyse(request):
    """POST (multipart : fichier) : chaque ligne contrôlée comme à l'import, rien n'est enregistré."""
    fichier = request.FILES.get("fichier")
    if not fichier:
        return _erreur("Choisissez le fichier Excel.")
    try:
        return Response(_importer(fichier, request.user, executer=False))
    except ErreurReprise as erreur:
        return _erreur(str(erreur))


@api_view(["POST"])
@authentication_classes(AUTHENTIFICATION)
@permission_classes([permissions.IsAuthenticated])
@parser_classes([MultiPartParser, FormParser, JSONParser])
def import_fichier(request):
    """POST (multipart : fichier) : lignes correctes reprises (un lot), lignes refusées signalées."""
    fichier = request.FILES.get("fichier")
    if not fichier:
        return _erreur("Choisissez le fichier Excel.")
    try:
        return Response(_importer(fichier, request.user, executer=True))
    except ErreurReprise as erreur:
        return _erreur(str(erreur))
