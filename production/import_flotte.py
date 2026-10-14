"""
Import Excel des véhicules d'une flotte automobile (devis flotte), sur le modèle du fichier remis par
OREOLE le 2026-10-09 (« FICHIER EXCEL FLOTTE.xlsx ») et de la liste NSIA « Liste des véhicules en
Automobile » qu'il reproduit :

    Catégorie 201  TARIF TOURISME                       ← une ligne par bloc de catégorie
    N° | Marque / Type | Puissance / CU. | Immatriculation / Energie / Zone | Valeur à neuf / Valeur
    vénale | Nombre de places | RC | CEDEAO | Protection Cond. | … | Vandalisme | Prime Nette
    1  | TOYOTA        | 8               | 2644EN01                         | 16 900 000   |   | 33 636 | …
       | VP            | 0               | Di                               | 1 200 000    | 5 | Oui    | …
       |               |                 | A                                |              |   | 4 000 000 000 …
    Total Catégorie / Total Générale                    ← ignorées

Un véhicule tient sur 1 à 3 lignes (prime / Oui-Non / capital par garantie, comme l'imprimé NSIA) ou
sur une seule ligne aux valeurs séparées par « / » (« TOYOTA / VP »). Les colonnes sont reconnues par
leur intitulé, dans n'importe quel ordre ; les colonnes de garanties absentes sont simplement ignorées.

Second format, le « FORMULAIRE ASSURANCE AUTO » d'OREOLE (compte rendu de la séance du 12-10-2026), rempli
par le client pour un ou plusieurs véhicules (un véhicule : devis mono, plusieurs : flotte) :

    FICHE POUR TARIFICATION - ASSURANCE AUTOMOBILE
    N° | Marque | Type de véhicule | N°Immatriculation | Date de 1ère mise en circulation | GENRE (cocher :
    Véhicule particulier, Camionnette, Camion, Engin de chantier, 02 Roues) | Energie | Puissance fiscale /
    Charge utile / Cylindrée | Nombre de places | Valeur neuve | Valeur vénale

Pas de garanties ni de catégorie : celles de l'écran s'appliquent (toutes les garanties de l'offre, que le
producteur décoche) ; valeur vénale du fichier comparée au tableau de dépréciation (CIEAMI).

Rien n'est écrit : l'analyse renvoie, pour chaque véhicule, les champs de l'écran Véhicule (référentiels
résolus : catégorie, marque, genre, carrosserie, usage, énergie) et l'offre de la catégorie qui couvre
le mieux les garanties marquées « Oui » sans celles marquées « Non » ; les garanties de l'offre marquées
« Non » sont à retirer du véhicule. L'écran ajoute ensuite ces véhicules à la flotte et les enregistre
par le chemin habituel (un sp_creation_devis par véhicule, puis sp_finalisation_devis).
"""

import re
import unicodedata
from datetime import date, datetime, time
from io import BytesIO

import openpyxl
from openpyxl.utils.datetime import from_excel, to_excel
from django.db import connection
from rest_framework import permissions, status
from rest_framework.authentication import BasicAuthentication
from rest_framework.decorators import api_view, authentication_classes, parser_classes, permission_classes
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.response import Response

from institutionnel.authentication import KnoxOrDemoTokenAuthentication

AUTHENTIFICATION = [KnoxOrDemoTokenAuthentication, BasicAuthentication]
TAILLE_MAXIMALE = 5 * 1024 * 1024

ID_PRODUIT_AUTO = 1

# Colonnes de garanties de l'imprimé NSIA (mêmes regroupements que les Conditions Particulières flotte,
# COLONNES_CP_FLOTTE de exportUtils.js) : clé, libellé, motif de l'intitulé (sans accent, minuscules),
# sous-garanties couvertes. Les motifs les plus précis d'abord (« recours express » avant « recours »).
COLONNES_GARANTIES = [
    ("recours_express", "Recours Express", r"express", [6, 173]),
    ("anticipe", "Remb. Anticipé", r"rembours|remb\b|anticip", [5]),
    ("defense", "Défense Recours", r"defense|^recours|^d\.?\s*r\.?$", [4]),
    ("rc", "RC", r"^r\.?\s*c\.?$|responsabilite", [1, 163, 165]),
    ("cedeao", "CEDEAO", r"cedea", [3]),
    ("protection", "Protection Cond.", r"protection|securite routiere|individuelle", [169, 16, 17, 18, 19, 20]),
    ("immobilisation", "Immobilisation", r"immob", [21]),
    ("collision", "Collision", r"collision", [8]),
    ("dommage", "Dommage", r"dommage", [7]),
    ("bris", "Bris de glace", r"bris", [9, 172]),
    ("assistance", "Assist'Car", r"assist", [199, 22, 147, 148, 174, 175, 177, 178, 179]),
    ("incendie", "Incendie", r"incendie", [10, 11, 164]),
    ("vol_main", "Vol à main armée", r"main", [13]),
    ("vol_accessoires", "Vol des accessoires", r"access", [15]),
    ("vandalisme", "Vandalisme", r"vandal", [14, 223]),
    ("vol_simple", "Vol simple", r"^vol", [12]),
]
GARANTIES_OBLIGATOIRES = {1, 3}
COLONNE_DE_GARANTIE = {sg: cle for cle, _, _, ids in COLONNES_GARANTIES for sg in ids}
LIBELLE_COLONNE = {cle: libelle for cle, libelle, _, _ in COLONNES_GARANTIES}

# Colonnes d'identification du véhicule
COLONNES_VEHICULE = [
    ("numero", r"^n\s*[°o]?\.?$|^numero$|^no$"),
    ("marque", r"marque"),
    ("puissance", r"puiss"),
    ("immatriculation", r"immatric"),
    ("valeurs", r"valeur"),
    ("places", r"place|nb\.?\s*pl"),
    ("prime_nette", r"prime|prine"),
]

OUI = {"oui", "o", "x", "yes", "y", "v"}
NON = {"non", "n", "no", "-"}

# Type du fichier (VP, CAMION NETTE…) → libellés des référentiels (genre, type de véhicule, carrosserie, usage)
TYPES_VEHICULE = [
    (r"camion\s*nette|camionnette|pick|bachee|plateau", ("Camionnette", "Véhicule Utilitaire", "CAMIONNETTE", "VEHICULE UTILITAIRE")),
    (r"fourgon", ("Fourgonnette", "Véhicule Utilitaire", "FOURGONNETTE", "VEHICULE UTILITAIRE")),
    (r"tracteur", ("Tracteur routier", "Véhicule Utilitaire", "TRACTEUR ROUTIER", "TRACTEUR ROUTIER")),
    (r"semi", ("Semi-remorque", "Véhicule Utilitaire", "SEMIE-REMORQUE", "SEMI-REMORQUE")),
    (r"remorque", ("Remorque", "Véhicule Utilitaire", "REMORQUE", "REMORQUE ATTELEE AU VEHICULE UTILITAIRE")),
    (r"camion", ("Camion", "Véhicule Utilitaire", "CAMION", "VEHICULE UTILITAIRE")),
    (r"mini\s*bus|minicar", ("Car", "Mini Car", "MINI-BUS", "VEHICULE MINIBUS")),
    (r"\bbus\b|\bcar\b|autocar", ("Car", "Auto Car", "BUS", "CAR")),
    (r"moto|scooter|cyclo|tricycle", ("Cyclomoteur (2/3 Roues)", "Cyclomoteur", "MOTOCYCLETTE", "SCOOTER, CYCLOMOTEURS, MOBILETTE, MOTOCYCLETTE")),
    (r"engin|chargeuse|bulldo|niveleuse", ("Engins de chantiers", "Véhicule Utilitaire", "ENGINS MOBILE DE CHANTIERS", "AUTRES ENGINS MOBILES DE CHANTIERS")),
]
TYPE_PAR_DEFAUT = ("Voiture (4 Roues)", "Véhicule Particulier", "CONDUITE INTERIEUR 4 PORTES", "VEHICULE DE TOURISME")

ENERGIES = [(r"^(di|die|dies|diesel|gasoil|go|gazole)", "dies"), (r"^(es|ess|essence|super|sp)", "essence"),
            (r"^(el|elec|electrique)", "electrique"), (r"^(hy|hyb|hybride)", "hybride")]
# Libellés de l'écran Véhicule (codes stdenergie : 1 Essence, 2 Diésel, 3 Électrique, 4 Hybride)
LIBELLES_ENERGIE_ECRAN = {"essence": "Essence", "dies": "Diesel", "electrique": "Électrique", "hybride": "Hybride"}


class ErreurImportFlotte(Exception):
    """Fichier illisible ou hors modèle : message affiché tel quel."""


def _norme(texte):
    texte = unicodedata.normalize("NFD", str(texte or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", texte.replace("\n", " ")).strip().lower()


def _vide(valeur):
    return valeur is None or str(valeur).strip() == ""


def _nombre(valeur):
    """Montant d'une cellule (nombre Excel ou texte « 16 900 000 ») ; None si ce n'en est pas un."""
    if isinstance(valeur, bool) or _vide(valeur):
        return None
    if isinstance(valeur, (int, float)):
        return int(round(valeur))
    texte = re.sub(r"[\s  ]", "", str(valeur)).replace(",", ".")
    if not re.fullmatch(r"-?\d+(\.\d+)?", texte):
        return None
    return int(round(float(texte)))


def _morceaux(valeur):
    """Valeurs d'une cellule « TOYOTA / VP » → ["TOYOTA", "VP"] (un nombre reste un nombre)."""
    if _vide(valeur):
        return []
    if isinstance(valeur, (int, float)) and not isinstance(valeur, bool):
        return [valeur]
    return [m.strip() for m in str(valeur).split("/") if m.strip()]


def _ressemble_immatriculation(valeur):
    texte = re.sub(r"[\s\-]", "", str(valeur or ""))
    return len(texte) >= 5 and len(re.findall(r"\d", texte)) >= 2 and bool(re.search(r"[A-Za-z]", texte))


# ------------------------------------------------------------------ lecture du fichier
def _colonnes(ligne):
    """Colonnes d'une ligne d'en-tête : {indice: clé} (identité du véhicule et garanties)."""
    trouvees, prises = {}, set()
    for indice, cellule in enumerate(ligne):
        intitule = _norme(cellule)
        if not intitule:
            continue
        for cle, motif in COLONNES_VEHICULE:
            if cle not in prises and re.search(motif, intitule):
                trouvees[indice] = cle
                prises.add(cle)
                break
        else:
            for cle, _, motif, _ in COLONNES_GARANTIES:
                if cle not in prises and re.search(motif, intitule):
                    trouvees[indice] = cle
                    prises.add(cle)
                    break
    return trouvees


def _est_entete(ligne):
    textes = [_norme(c) for c in ligne if not _vide(c)]
    return any("marque" in t for t in textes) and any("immatric" in t for t in textes)


def _categorie(ligne):
    """« Catégorie 201  TARIF TOURISME » → ("201", "TARIF TOURISME") ; ("", "") pour « Catégorie XXX »."""
    texte = " ".join(str(c) for c in ligne if not _vide(c))
    if not _norme(texte).startswith("categorie"):
        return None
    reste = re.sub(r"^\S+\s*", "", texte.strip())
    code = re.match(r"(\d{3})\b\s*(.*)", reste)
    if not code:
        return ("", "")
    libelle = re.sub(r"\bX{2,}\b", "", code.group(2)).strip(" -:")
    return (code.group(1), " ".join(libelle.split()))


def _cellule(valeur):
    """Le modèle a des cellules au format date : « 7 » y est lu 07/01/1900. Aucune colonne n'attend de
    date : on revient au nombre saisi (numéro de série Excel)."""
    if isinstance(valeur, (datetime, date, time)):
        serie = to_excel(valeur)
        return int(serie) if float(serie).is_integer() else serie
    return valeur


def lire_fichier(contenu):
    """Blocs de lignes par véhicule : [{ligne, categorie, cellules: {clé: [(rang, valeur), …]}}]."""
    try:
        classeur = openpyxl.load_workbook(BytesIO(contenu), data_only=True, read_only=True)
    except Exception:  # fichier corrompu, .xls, autre format
        raise ErreurImportFlotte("Fichier illisible : enregistrez la liste des véhicules au format Excel .xlsx.")
    vehicules, colonnes, categorie = [], None, ("", "")
    courant = None
    for feuille in classeur.worksheets:
        for numero, ligne in enumerate(feuille.iter_rows(values_only=True), start=1):
            ligne = [_cellule(c) for c in ligne]
            if all(_vide(c) for c in ligne):
                continue
            cat = _categorie(ligne)
            if cat is not None:
                categorie, courant = cat, None
                continue
            premier = next((_norme(c) for c in ligne if not _vide(c)), "")
            if premier.startswith("total"):
                courant = None
                continue
            if _est_entete(ligne):
                colonnes, courant = _colonnes(ligne), None
                continue
            if not colonnes:
                continue
            valeurs = {cle: ligne[i] for i, cle in colonnes.items() if i < len(ligne) and not _vide(ligne[i])}
            if not valeurs:
                continue
            immat = valeurs.get("immatriculation")
            nouveau = (
                courant is None
                or "numero" in valeurs
                or (_ressemble_immatriculation(_morceaux(immat)[0] if _morceaux(immat) else "") and courant["immat_vue"])
                or courant["rangs"] >= 3
            )
            if nouveau:
                courant = {"ligne": numero, "categorie": categorie, "cellules": {}, "rangs": 0, "immat_vue": False}
                vehicules.append(courant)
            rang = courant["rangs"]
            for cle, valeur in valeurs.items():
                courant["cellules"].setdefault(cle, []).append((rang, valeur))
            if immat is not None and _ressemble_immatriculation(_morceaux(immat)[0] if _morceaux(immat) else ""):
                courant["immat_vue"] = True
            courant["rangs"] += 1
    classeur.close()
    if colonnes is None:
        raise ErreurImportFlotte(
            "Ligne d'en-tête introuvable : le fichier doit reprendre les colonnes du modèle "
            "(N°, Marque / Type, Puissance / CU., Immatriculation / Energie / Zone, Valeur à neuf / Valeur vénale…)."
        )
    return vehicules


def _parties(cellules, cle, nombre):
    """Valeurs d'une colonne d'identité par rang : une ligne par valeur (imprimé NSIA, « Marque » puis
    « Type » en dessous) ou plusieurs valeurs séparées par « / » dans une même cellule (« TOYOTA / VP »)."""
    resultat = [None] * nombre
    for rang, valeur in cellules.get(cle, []):
        morceaux = _morceaux(valeur)
        if len(morceaux) > 1:
            for i, morceau in enumerate(morceaux[:nombre]):
                if resultat[i] is None:
                    resultat[i] = morceau
        elif morceaux and rang < nombre and resultat[rang] is None:
            resultat[rang] = morceaux[0]
    return resultat


def _premier_nombre(cellules, cle):
    return next((_nombre(m) for _, v in cellules.get(cle, []) for m in _morceaux(v) if _nombre(m) is not None), None)


def _garanties_fichier(cellules, rangs):
    """Par colonne de garantie : souscrite (True/False/None), prime et capital lus dans le fichier."""
    garanties = {}
    for cle, _, _, _ in COLONNES_GARANTIES:
        if cle not in cellules:
            continue
        souscrite, prime, capital = None, None, None
        for rang, valeur in cellules[cle]:
            for morceau in _morceaux(valeur):
                mot = _norme(morceau)
                if mot in OUI:
                    souscrite = True
                elif mot in NON:
                    souscrite = False
                else:
                    montant = _nombre(morceau)
                    if montant is None:
                        continue
                    # Imprimé NSIA : prime sur la 1re ligne, capital sur la 3e
                    if rang >= 2 and rangs >= 3:
                        capital = montant
                    elif prime is None:
                        prime = montant
                    else:
                        capital = montant
        if souscrite is None and prime:
            souscrite = True
        garanties[cle] = {"Colonne": cle, "Libelle": LIBELLE_COLONNE[cle], "Souscrite": souscrite,
                          "Prime": prime, "Capital": capital}
    return garanties


# ------------------------------------------------------------------ référentiels
def _referentiels(cursor, id_compagnie):
    def table(sql):
        cursor.execute(sql)
        return cursor.fetchall()

    tarifs = table("SELECT idtarif, libelle, COALESCE(codecategorie, '') FROM stdtarif WHERE codecategorie LIKE '2%%' ORDER BY idtarif")
    cursor.execute(
        """
        SELECT og.idoffre, og.idsousgarantie FROM stdoffregarantie og WHERE og.idcompagnie = %s
        """,
        [id_compagnie],
    )
    garanties_offre = {}
    for id_offre, id_sg in cursor.fetchall():
        garanties_offre.setdefault(id_offre, set()).add(id_sg)
    cursor.execute(
        "SELECT codeformule, libellelongformule FROM fn_formule_securite_routiere_compagnie(%s)", [id_compagnie]
    )
    formules = cursor.fetchall()
    cursor.execute(
        "SELECT idoption, libelleoption, primenette FROM stdassistanceautomobile WHERE idcompagnie IN (%s, 0) ORDER BY idoption",
        [id_compagnie],
    )
    assistances = cursor.fetchall()
    return {
        "tarifs": tarifs,
        "marques": table("SELECT idmarque, libellemarque FROM stdmarque WHERE idmarque <> 0"),
        "genres": table("SELECT idgenre, libellegenre FROM stdgenrevehicule WHERE idgenre <> 0"),
        "types": table("SELECT id, libelletype FROM stdtypevehicule WHERE id <> 0"),
        "carrosseries": table("SELECT idcarrosserie, libellecarrosserie FROM stdcarrosserie"),
        "usages": table("SELECT idusage, libelleusage FROM stdusage WHERE idusage <> 0"),
        "energies": table("SELECT idenergie, libelle FROM stdenergie WHERE idenergie <> 0"),
        "garanties_offre": garanties_offre,
        "formules": formules,
        "assistances": assistances,
        "offres_tarif": {},
    }


def _par_libelle(lignes, libelle, debut=False):
    cible = _norme(libelle)
    for identifiant, nom in lignes:
        if _norme(nom) == cible or (debut and _norme(nom).startswith(cible)):
            return identifiant, nom
    return None, None


def _marque(ref, texte):
    compacte = re.sub(r"[^a-z0-9]", "", _norme(texte))
    if not compacte:
        return None, None
    for identifiant, nom in ref["marques"]:
        if re.sub(r"[^a-z0-9]", "", _norme(nom)) == compacte:
            return identifiant, nom
    # « MERCEDES » pour « MERCEDES-BENZ », « LAND » pour « LAND ROVER » : seulement si un seul candidat
    candidats = [(i, n) for i, n in ref["marques"] if re.sub(r"[^a-z0-9]", "", _norme(n)).startswith(compacte)]
    return candidats[0] if len(candidats) == 1 else (None, None)


def _tarif(ref, code, libelle, id_tarif_defaut):
    """Tarif de la catégorie du bloc : celui dont le libellé correspond, sinon le tarif de base (plus petit id)."""
    if not code:
        return next(((i, n, c) for i, n, c in ref["tarifs"] if i == id_tarif_defaut), (None, None, None))
    candidats = [(i, n, c) for i, n, c in ref["tarifs"] if c == code]
    if not candidats:
        return None, None, code
    cible = _norme(libelle)
    for i, n, c in candidats:
        if cible and _norme(n) == cible:
            return i, n, c
    # Formule nommée (« TARIF EBENE PRISME ») : tous ses mots distinctifs figurent dans le libellé du
    # fichier ; sinon le tarif de base de la catégorie (« TARIF TOURISME » → PROMENADE ET AFFAIRES)
    fichier = _mots_distinctifs(libelle)
    nommes = [(len(_mots_distinctifs(n)), -i, (i, n, c)) for i, n, c in candidats
              if _mots_distinctifs(n) and _mots_distinctifs(n) <= fichier]
    return max(nommes)[2] if nommes else candidats[0]


MOTS_GENERIQUES = {"tarif", "pour", "compte", "assure", "assuree", "transport", "promenade", "affaires", "affaire",
                   "personne", "morale", "vehicule", "vehicules", "categorie", "tourisme", "les", "des"}


def _mots_distinctifs(texte):
    return set(re.findall(r"[a-z]{3,}", _norme(texte))) - MOTS_GENERIQUES


def _offres_tarif(cursor, ref, id_tarif):
    if id_tarif not in ref["offres_tarif"]:
        cursor.execute("SELECT id_offre, libelle_offre FROM fn_liste_offre_produit(%s, %s)", [ID_PRODUIT_AUTO, id_tarif])
        ref["offres_tarif"][id_tarif] = [(i, n) for i, n in cursor.fetchall() if ref["garanties_offre"].get(i)]
    return ref["offres_tarif"][id_tarif]


def _choisir_offre(offres, ref, demandees, refusees, id_offre_defaut):
    """Offre couvrant toutes les colonnes « Oui » (sinon le plus possible), avec le moins de colonnes « Non »."""
    if not offres:
        return None, None, set()
    if not demandees and not refusees:
        choix = next(((i, n) for i, n in offres if i == id_offre_defaut), offres[0])
        return choix[0], choix[1], set()

    def colonnes(id_offre):
        return {COLONNE_DE_GARANTIE[g] for g in ref["garanties_offre"][id_offre] if g in COLONNE_DE_GARANTIE}

    def score(offre):
        couvertes = colonnes(offre[0])
        return (len(demandees - couvertes), len(couvertes & refusees), len(ref["garanties_offre"][offre[0]]),
                offre[0] != id_offre_defaut)

    meilleure = min(offres, key=score)
    return meilleure[0], meilleure[1], demandees - colonnes(meilleure[0])


def _formule_securite(ref, capital):
    """Formule de sécurité routière dont le capital décès est celui du fichier, sinon la première."""
    if not ref["formules"]:
        return None
    if capital:
        for code, libelle in ref["formules"]:
            deces = re.search(r"d[ée]c[èe]s\s*:\s*(\d+)", libelle or "", re.I)
            if deces and int(deces.group(1)) == capital:
                return code
    return sorted(ref["formules"], key=lambda f: (len(f[0]), f[0]))[0][0]


def _option_assistance(ref, prime):
    options = [(i, n, int(p or 0)) for i, n, p in ref["assistances"] if int(p or 0) > 0]
    if not options:
        return None
    if prime:
        for i, n, p in options:
            if p == prime:
                return i
    return options[0][0]


# ------------------------------------------------------------------ analyse
def analyser_flotte(contenu, id_compagnie, id_tarif_defaut=None, id_offre_defaut=None):
    blocs = lire_fichier(contenu)
    vehicules, avertissements = [], []
    with connection.cursor() as cursor:
        ref = _referentiels(cursor, id_compagnie)
        vues = set()
        for bloc in blocs:
            c = bloc["cellules"]
            alertes, erreurs = [], []
            marque_txt, type_txt = (str(x or "").strip() for x in _parties(c, "marque", 2))
            puissance, charge = _parties(c, "puissance", 2)
            immat, energie_txt, zone = (str(x or "").strip() for x in _parties(c, "immatriculation", 3))
            neuve, venale = _parties(c, "valeurs", 2)
            places = _premier_nombre(c, "places")
            prime_nette = _premier_nombre(c, "prime_nette")
            numero = str(_parties(c, "numero", 1)[0] or "")

            immatriculation = re.sub(r"\s+", "", str(immat or "")).upper()
            if not immatriculation:
                erreurs.append("Immatriculation absente.")
            elif immatriculation in vues:
                erreurs.append(f"Immatriculation {immatriculation} en double dans le fichier.")
            if immatriculation:
                vues.add(immatriculation)

            code_cat, libelle_cat = bloc["categorie"]
            id_tarif, libelle_tarif, code_tarif = _tarif(ref, code_cat, libelle_cat, id_tarif_defaut)
            if not id_tarif:
                erreurs.append(f"Catégorie {code_cat} inconnue des tarifs Automobile." if code_cat
                               else "Catégorie non précisée dans le fichier.")
            elif not code_cat:
                alertes.append("Catégorie non précisée dans le fichier : catégorie de l'écran retenue.")

            id_marque, libelle_marque = _marque(ref, marque_txt)
            if not marque_txt:
                erreurs.append("Marque absente.")
            elif not id_marque:
                erreurs.append(f"Marque « {marque_txt} » inconnue du référentiel.")

            genre, type_vehicule, carrosserie, usage = next(
                (libelles for motif, libelles in TYPES_VEHICULE if re.search(motif, _norme(type_txt))), TYPE_PAR_DEFAUT)
            id_genre, genre = _par_libelle(ref["genres"], genre)
            id_type, type_vehicule = _par_libelle(ref["types"], type_vehicule, debut=True)
            id_carrosserie, carrosserie = _par_libelle(ref["carrosseries"], carrosserie)
            id_usage, usage = _par_libelle(ref["usages"], usage)

            energie_cle = next((cle for motif, cle in ENERGIES if re.search(motif, _norme(energie_txt))), None)
            id_energie = next((i for i, n in ref["energies"] if energie_cle and _norme(n).startswith(energie_cle[:4])), None)
            if not energie_cle:
                alertes.append(f"Énergie « {energie_txt} » non reconnue : Essence retenue." if energie_txt
                               else "Énergie absente : Essence retenue.")
                energie_cle, id_energie = "essence", 1

            puissance, charge, neuve, venale = (_nombre(puissance), _nombre(charge), _nombre(neuve), _nombre(venale))
            if not puissance:
                alertes.append("Puissance fiscale absente.")
            if not places:
                alertes.append("Nombre de places absent : 5 retenu.")
                places = 5
            if neuve and venale and venale > neuve:
                alertes.append("Valeur vénale supérieure à la valeur à neuf.")

            garanties = _garanties_fichier(c, bloc["rangs"])
            demandees = {k for k, g in garanties.items() if g["Souscrite"] is True}
            refusees = {k for k, g in garanties.items() if g["Souscrite"] is False}
            accessoire = (garanties.get("vol_accessoires") or {}).get("Capital") or 0
            if accessoire and venale is not None and accessoire > venale:
                alertes.append("Capital « Vol des accessoires » supérieur à la valeur vénale : non repris.")
                accessoire = 0

            id_offre, libelle_offre, manquantes = (None, None, set())
            if id_tarif:
                offres = _offres_tarif(cursor, ref, id_tarif)
                id_offre, libelle_offre, manquantes = _choisir_offre(offres, ref, demandees, refusees, id_offre_defaut)
                if not id_offre:
                    erreurs.append(f"Aucune offre de la catégorie {libelle_tarif} n'est paramétrée pour cette compagnie.")
            # RC et CEDEAO restent (obligatoires, comme dans « Garanties de l'offre »)
            retirees = sorted(g for g in ref["garanties_offre"].get(id_offre, set())
                              if COLONNE_DE_GARANTIE.get(g) in refusees and g not in GARANTIES_OBLIGATOIRES)
            if manquantes:
                alertes.append("Garanties demandées absentes des offres de la catégorie : "
                               + ", ".join(LIBELLE_COLONNE[k] for k in sorted(manquantes))
                               + " (à ajouter dans « Garanties de l'offre »).")

            formule = "AUCUNE"
            if "protection" in demandees:
                formule = _formule_securite(ref, (garanties["protection"] or {}).get("Capital")) or "AUCUNE"
                if formule == "AUCUNE":
                    alertes.append("Protection du conducteur demandée : aucune formule de sécurité routière pour cette compagnie.")
            option = 0
            if "assistance" in demandees:
                option = _option_assistance(ref, garanties["assistance"]["Prime"]) or 0
                if not option:
                    alertes.append("Assist'Car demandée : aucune option d'assistance pour cette compagnie.")

            vehicules.append({
                "Ligne": bloc["ligne"],
                "Numero": numero,
                "Erreurs": erreurs,
                "Avertissements": alertes,
                "Fichier": {
                    "Categorie": " ".join(x for x in (code_cat, libelle_cat) if x), "Marque": marque_txt, "Type": type_txt,
                    "Energie": energie_txt, "Zone": zone, "PrimeNette": prime_nette,
                    "Garanties": list(garanties.values()),
                },
                # Champs de l'écran Véhicule (instantaneVehicule de NewAutoQuotePage)
                "Vehicule": {
                    "idTarif": id_tarif, "categorieContrat": libelle_tarif or "", "codeCategorie": code_tarif or code_cat,
                    "codeUsage": id_usage or 1, "usageVehicule": usage or "",
                    "idCarrosserie": id_carrosserie or 1, "carrosserie": carrosserie or "",
                    "idOffre": id_offre, "offreSelectionnee": libelle_offre or "",
                    "codeCarburant": id_energie or 1, "energie": LIBELLES_ENERGIE_ECRAN[energie_cle],
                    "codeFormuleSecuriteRoutiere": formule, "idOptionAssistance": option,
                    "idMarque": id_marque, "marqueVehicule": libelle_marque or marque_txt,
                    "idGenreVehicule": id_genre or 0, "genreVehicule": genre or "",
                    "idTypeVehicule": id_type or 0, "typeCommercial": type_vehicule or "",
                    "modeleVehicule": str(type_txt or "").upper(),
                    "puissanceFiscale": puissance or 0, "chargeUtile": charge or 0, "nombrePlace": places,
                    "valeurNeuf": neuve or 0, "valeurVenale": venale or 0, "valeurAccessoire": accessoire,
                    "immatriculation": immatriculation,
                },
                "GarantiesRetirees": retirees,
                "ColonnesRetirees": sorted({LIBELLE_COLONNE[COLONNE_DE_GARANTIE[g]] for g in retirees}),
            })
    if not vehicules:
        avertissements.append("Aucun véhicule trouvé sous la ligne d'en-tête du fichier.")
    return {
        "Vehicules": vehicules,
        "Avertissements": avertissements,
        "Valides": sum(1 for v in vehicules if not v["Erreurs"]),
    }


# ------------------------------------------------------------------ formulaire OREOLE (mono ou flotte)
COLONNES_FORMULAIRE = [
    ("numero", r"^n\s*[°o]?\.?$|^numero$|^no$"),
    ("marque", r"marque"),
    ("type", r"^type"),
    ("immatriculation", r"immatric"),
    ("date_mec", r"circulation|m\.?\s*e\.?\s*c\.?"),
    ("energie", r"energie|carburant"),
    ("puissance", r"puissance|charge|cylindr"),
    ("places", r"place"),
    ("valeur_neuve", r"valeur\s*(neuve|a neuf|catalogue)"),
    ("valeur_venale", r"venale"),
]
# Colonnes « GENRE (cocher la bonne case) » → libellés des référentiels (genre, type, carrosserie, usage)
GENRES_FORMULAIRE = [
    (r"camion\s*nette|camionnette", "Camionnette", ("Camionnette", "Véhicule Utilitaire", "CAMIONNETTE", "VEHICULE UTILITAIRE")),
    (r"camion", "Camion", ("Camion", "Véhicule Utilitaire", "CAMION", "VEHICULE UTILITAIRE")),
    (r"engin", "Engin de chantier", ("Engins de chantiers", "Véhicule Utilitaire", "ENGINS MOBILE DE CHANTIERS", "AUTRES ENGINS MOBILES DE CHANTIERS")),
    (r"roue|cyclo|moto", "02 roues", ("Cyclomoteur (2/3 Roues)", "Cyclomoteur", "MOTOCYCLETTE", "SCOOTER, CYCLOMOTEURS, MOBILETTE, MOTOCYCLETTE")),
    (r"particulier|tourisme", "Véhicule particulier", TYPE_PAR_DEFAUT),
]


def _date_formulaire(valeur):
    """Date de mise en circulation : cellule date, « 01/05/2018 », « 2018-05-01 » ou année seule (1er janvier)."""
    if isinstance(valeur, datetime):
        return valeur.date(), None
    if isinstance(valeur, date):
        return valeur, None
    if isinstance(valeur, (int, float)) and not isinstance(valeur, bool):
        if 1950 <= valeur <= 2100:
            return date(int(valeur), 1, 1), "Année de mise en circulation seule : 1er janvier retenu."
        try:  # numéro de série Excel
            return from_excel(valeur).date(), None
        except Exception:  # noqa: BLE001
            return None, None
    texte = str(valeur or "").strip()
    for format_ in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y", "%d.%m.%Y"):
        try:
            return datetime.strptime(texte[:10], format_).date(), None
        except ValueError:
            continue
    annee = re.fullmatch(r"(19[5-9]\d|20\d\d)", texte)
    if annee:
        return date(int(annee.group(1)), 1, 1), "Année de mise en circulation seule : 1er janvier retenu."
    return None, None


def _puissance_formulaire(valeur, genre):
    """« Puissance fiscale (CV) / Charge utile (kg ou t) / Cylindrée (cm³) » → (puissance, charge kg, cylindrée, alerte)."""
    puissance = charge = cylindree = None
    alerte = None
    for morceau in _morceaux(valeur):
        texte = _norme(morceau)
        chiffres = re.sub(r"[^\d.,]", "", texte).replace(",", ".")
        if not re.fullmatch(r"\d+(\.\d+)?", chiffres or ""):
            continue
        nombre = float(chiffres)
        if re.search(r"cm|cc", texte):
            cylindree = int(round(nombre))
        elif "kg" in texte:
            charge = int(round(nombre))
        elif re.search(r"tonne|\d\s*t$|\bt\b", texte):
            charge = int(round(nombre * 1000))
        elif "cv" in texte:
            puissance = int(round(nombre))
        elif genre in ("Camion", "Camionnette") and nombre > 100:
            charge = int(round(nombre))
        elif genre == "02 roues" and nombre > 40:
            cylindree = int(round(nombre))
        elif puissance is None:
            puissance = int(round(nombre))
            if genre in ("Camion", "Camionnette"):
                alerte = "Puissance lue sans unité : charge utile du camion à vérifier."
    return puissance, charge, cylindree, alerte


def lire_formulaire(contenu):
    """Lignes du formulaire OREOLE : [{ligne, valeurs: {clé: valeur}, genres}], ou None s'il n'est pas à ce format."""
    try:
        classeur = openpyxl.load_workbook(BytesIO(contenu), data_only=True, read_only=True)
    except Exception:  # noqa: BLE001
        raise ErreurImportFlotte("Fichier illisible : enregistrez le formulaire au format Excel .xlsx.")
    lignes = []
    trouve = False
    for feuille in classeur.worksheets:
        colonnes, genres, attente_genres = None, {}, False
        for numero, ligne in enumerate(feuille.iter_rows(values_only=True), start=1):
            ligne = list(ligne)
            textes = [_norme(c) for c in ligne]
            if colonnes is None:
                if any("immatric" in t for t in textes) and any(re.search(r"circulation|m\.?\s*e\.?\s*c", t) for t in textes):
                    colonnes, prises = {}, set()
                    for i, t in enumerate(textes):
                        if not t:
                            continue
                        for cle, motif in COLONNES_FORMULAIRE:
                            if cle not in prises and re.search(motif, t):
                                colonnes[i] = cle
                                prises.add(cle)
                                break
                    attente_genres = True
                    trouve = True
                continue
            if attente_genres and any(re.search(r"particulier|camion|engin|roue", t) for t in textes):
                for i, t in enumerate(textes):
                    for motif, libelle, _ in GENRES_FORMULAIRE:
                        if t and re.search(motif, t):
                            genres[i] = libelle
                            break
                attente_genres = False
                continue
            attente_genres = False
            valeurs = {cle: ligne[i] for i, cle in colonnes.items() if i < len(ligne) and not _vide(ligne[i])}
            coches = [libelle for i, libelle in genres.items() if i < len(ligne) and not _vide(ligne[i])]
            if not ({"marque", "immatriculation", "type", "valeur_neuve"} & set(valeurs)) and not coches:
                continue  # ligne numérotée vide du modèle
            lignes.append({"ligne": numero, "valeurs": valeurs, "genres": coches})
    classeur.close()
    return lignes if trouve else None


def analyser_formulaire(lignes, id_compagnie, id_tarif_defaut=None, id_offre_defaut=None, date_effet=None):
    from configuration_api.parametrage_auto import calculer_valeurs_venales

    vehicules = []
    with connection.cursor() as cursor:
        ref = _referentiels(cursor, id_compagnie)
        id_tarif, libelle_tarif, code_tarif = next(((i, n, c) for i, n, c in ref["tarifs"] if i == id_tarif_defaut), (None, None, None))
        offres = _offres_tarif(cursor, ref, id_tarif) if id_tarif else []
        id_offre, libelle_offre = next(((i, n) for i, n in offres if i == id_offre_defaut), offres[0] if offres else (None, None))
        vues = set()
        for l in lignes:
            v = l["valeurs"]
            alertes, erreurs = [], []
            marque_txt = str(v.get("marque") or "").strip()
            type_txt = str(v.get("type") or "").strip()
            immatriculation = re.sub(r"\s+", "", str(v.get("immatriculation") or "")).upper()
            if not immatriculation:
                erreurs.append("Immatriculation absente.")
            elif immatriculation in vues:
                erreurs.append(f"Immatriculation {immatriculation} en double dans le fichier.")
            vues.add(immatriculation)
            if not id_tarif:
                erreurs.append("Choisissez la catégorie du véhicule à l'écran avant d'importer le formulaire (il ne la donne pas).")
            elif not id_offre:
                erreurs.append(f"Aucune offre de la catégorie {libelle_tarif} n'est paramétrée pour cette compagnie.")

            id_marque, libelle_marque = _marque(ref, marque_txt)
            if not marque_txt:
                erreurs.append("Marque absente.")
            elif not id_marque:
                erreurs.append(f"Marque « {marque_txt} » inconnue du référentiel.")

            if len(l["genres"]) > 1:
                alertes.append(f"Plusieurs genres cochés ({', '.join(l['genres'])}) : {l['genres'][0]} retenu.")
            genre_coche = l["genres"][0] if l["genres"] else None
            if genre_coche:
                libelles = next(lib for _, nom, lib in GENRES_FORMULAIRE if nom == genre_coche)
            else:
                libelles = next((lib for motif, lib in TYPES_VEHICULE if re.search(motif, _norme(type_txt))), TYPE_PAR_DEFAUT)
                alertes.append("Genre non coché : déduit du type de véhicule." if type_txt else "Genre non coché : véhicule particulier retenu.")
            genre, type_vehicule, carrosserie, usage = libelles
            id_genre, genre = _par_libelle(ref["genres"], genre)
            id_type, type_vehicule = _par_libelle(ref["types"], type_vehicule, debut=True)
            id_carrosserie, carrosserie = _par_libelle(ref["carrosseries"], carrosserie)
            id_usage, usage = _par_libelle(ref["usages"], usage)
            # Le formulaire ne donne pas la catégorie : celle de l'écran, à vérifier si le genre ne lui correspond pas
            if id_tarif and ((genre_coche in ("Camion", "Camionnette", "Engin de chantier") and code_tarif in ("201", "212"))
                             or (genre_coche == "02 roues" and code_tarif != "205")
                             or (genre_coche == "Engin de chantier" and code_tarif != "209")):
                alertes.append(f"{genre_coche} importé dans la catégorie {code_tarif} {libelle_tarif} : catégorie à vérifier.")

            energie_txt = str(v.get("energie") or "").strip()
            energie_cle = next((cle for motif, cle in ENERGIES if re.search(motif, _norme(energie_txt))), None)
            id_energie = next((i for i, n in ref["energies"] if energie_cle and _norme(n).startswith(energie_cle[:4])), None)
            if not energie_cle:
                alertes.append(f"Énergie « {energie_txt} » non reconnue : Essence retenue." if energie_txt else "Énergie absente : Essence retenue.")
                energie_cle, id_energie = "essence", 1

            date_mec, alerte_date = _date_formulaire(v.get("date_mec"))
            if alerte_date:
                alertes.append(alerte_date)
            if not date_mec:
                alertes.append("Date de 1ère mise en circulation absente ou illisible : valeur vénale non contrôlée.")
            puissance, charge, cylindree, alerte_puissance = _puissance_formulaire(v.get("puissance"), genre_coche)
            if alerte_puissance:
                alertes.append(alerte_puissance)
            if not puissance and genre_coche != "02 roues":
                alertes.append("Puissance fiscale absente.")
            places = _nombre(v.get("places"))
            if not places:
                places = 2 if genre_coche == "02 roues" else 5
                alertes.append(f"Nombre de places absent : {places} retenu.")
            neuve, venale_fichier = _nombre(v.get("valeur_neuve")), _nombre(v.get("valeur_venale"))
            if not neuve:
                alertes.append("Valeur neuve absente.")

            tableau = {}
            if neuve and date_mec:
                tableau = calculer_valeurs_venales([{
                    "IdGenre": id_genre, "IdTypeVehicule": id_type, "IdCarrosserie": id_carrosserie, "IdTarif": id_tarif,
                    "ChargeUtile": charge or 0, "ValeurNeuve": neuve, "DateMec": date_mec}], date_effet)[0]
            venale_tableau = tableau.get("ValeurVenale")
            venale = venale_fichier or venale_tableau or 0
            montant = lambda x: f"{x:,} F".replace(",", " ")  # noqa: E731
            if venale_fichier and venale_tableau and abs(venale_fichier - venale_tableau) > max(venale_tableau * 0.01, 1000):
                alertes.append(f"Valeur vénale du fichier ({montant(venale_fichier)}) différente du tableau de dépréciation "
                               f"({montant(venale_tableau)} : {tableau.get('Explication')}).")
            elif not venale_fichier and venale_tableau:
                alertes.append(f"Valeur vénale absente du fichier : tableau de dépréciation appliqué ({tableau.get('Explication')}).")
            elif not venale_fichier and neuve and date_mec and venale_tableau is None:
                alertes.append(f"Valeur vénale absente : {tableau.get('Explication') or 'tableau de dépréciation non applicable'}, à saisir.")
            if neuve and venale > neuve:
                alertes.append("Valeur vénale supérieure à la valeur neuve.")

            vehicules.append({
                "Ligne": l["ligne"],
                "Numero": str(v.get("numero") or ""),
                "Erreurs": erreurs,
                "Avertissements": alertes,
                "Fichier": {
                    "Categorie": "", "Marque": marque_txt, "Type": type_txt, "Energie": energie_txt, "Zone": "",
                    "PrimeNette": None, "Garanties": [], "Genre": genre_coche or "",
                    "DateMec": date_mec.isoformat() if date_mec else None, "Cylindree": cylindree,
                    "ValeurVenale": venale_fichier, "ValeurVenaleTableau": venale_tableau,
                    "ExplicationTableau": tableau.get("Explication"),
                },
                "Vehicule": {
                    "idTarif": id_tarif, "categorieContrat": libelle_tarif or "", "codeCategorie": code_tarif or "",
                    "codeUsage": id_usage or 1, "usageVehicule": usage or "",
                    "idCarrosserie": id_carrosserie or 1, "carrosserie": carrosserie or "",
                    "idOffre": id_offre, "offreSelectionnee": libelle_offre or "",
                    "codeCarburant": id_energie or 1, "energie": LIBELLES_ENERGIE_ECRAN[energie_cle],
                    "codeFormuleSecuriteRoutiere": "AUCUNE", "idOptionAssistance": 0,
                    "idMarque": id_marque, "marqueVehicule": libelle_marque or marque_txt,
                    "idGenreVehicule": id_genre or 0, "genreVehicule": genre or "",
                    "idTypeVehicule": id_type or 0, "typeCommercial": type_vehicule or "",
                    "modeleVehicule": type_txt.upper(),
                    "puissanceFiscale": puissance or 0, "chargeUtile": charge or 0, "nombrePlace": places,
                    "valeurNeuf": neuve or 0, "valeurVenale": venale, "valeurAccessoire": 0,
                    "immatriculation": immatriculation,
                    **({"dateMec": date_mec.isoformat()} if date_mec else {}),
                },
                "GarantiesRetirees": [],
                "ColonnesRetirees": [],
            })
    return {
        "Format": "formulaire",
        "Vehicules": vehicules,
        "Avertissements": [] if vehicules else ["Aucun véhicule renseigné dans le formulaire."],
        "Valides": sum(1 for x in vehicules if not x["Erreurs"]),
    }


# ------------------------------------------------------------------ route
@api_view(["POST"])
@authentication_classes(AUTHENTIFICATION)
@permission_classes([permissions.IsAuthenticated])
@parser_classes([MultiPartParser, FormParser])
def analyser_import_flotte(request):
    """POST /api/flotte/import/analyse/ (fichier_excel, IdCompagnie, IdTarif, IdOffre et DateEffet de l'écran) —
    n'écrit rien. Formulaire OREOLE (un ou plusieurs véhicules) ou liste NSIA des véhicules d'une flotte."""
    fichier = request.FILES.get("fichier_excel")
    if not fichier:
        return Response({"error": "Sélectionnez le fichier Excel des véhicules de la flotte."}, status=status.HTTP_400_BAD_REQUEST)
    if fichier.size > TAILLE_MAXIMALE:
        return Response({"error": "Fichier trop volumineux (5 Mo au plus)."}, status=status.HTTP_400_BAD_REQUEST)

    def entier(cle):
        try:
            return int(request.data.get(cle) or 0) or None
        except (TypeError, ValueError):
            return None

    if not entier("IdCompagnie"):
        return Response({"error": "Choisissez la compagnie du devis avant d'importer les véhicules."}, status=status.HTTP_400_BAD_REQUEST)
    try:
        contenu = fichier.read()
        formulaire = lire_formulaire(contenu)
        if formulaire is not None:
            return Response(analyser_formulaire(formulaire, entier("IdCompagnie"), entier("IdTarif"), entier("IdOffre"),
                                                request.data.get("DateEffet") or None))
        return Response({"Format": "liste", **analyser_flotte(contenu, entier("IdCompagnie"), entier("IdTarif"), entier("IdOffre"))})
    except ErreurImportFlotte as erreur:
        return Response({"error": str(erreur)}, status=status.HTTP_400_BAD_REQUEST)
