import io
import re
import unicodedata
from copy import copy
from pathlib import Path

from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework import status
from openpyxl import load_workbook

from configuration_api.models import Tarif

from .models import Contrat, Devis, DevisDetail


MODELE_DEMANDE_ATTESTATION = (
    Path(__file__).resolve().parent
    / "templates"
    / "modele_demande_edition_attestation_auto.xlsx"
)


def _nom_client(client):
    return " ".join(part for part in (client.Nom, client.Prenoms) if part).strip()


def _date_texte(valeur):
    return valeur.strftime("%Y-%m-%d") if valeur else ""


def _prime_rc(detail):
    for garantie in detail.garanties.all():
        if not garantie.Acquise:
            continue
        libelle = unicodedata.normalize(
            "NFD", garantie.IdGarantie.LibelleSousGarantie or ""
        )
        libelle = "".join(caractere for caractere in libelle if not unicodedata.combining(caractere)).lower()
        if "responsabilite civile" in libelle or libelle.strip() in {"rc", "rco", "rca"}:
            return garantie.PrimeNette
    return None


def _ligne_attestation(devis, detail, contrat, souscripteur, assure, tarifs):
    tarif = tarifs.get(detail.idtarif)
    categorie = ""
    if tarif:
        categorie = tarif.code_categorie_asaci or (
            tarif.IdCategorie.code_asaci if tarif.IdCategorie else ""
        ) or tarif.CodeCategorie

    type_souscripteur = getattr(souscripteur.idtypeclient, "code_type", "")
    energie = getattr(detail.essence, "CodeEnergie", "")
    genre = getattr(detail.idgenrevehicule, "CodeGenre", "")
    marque = getattr(detail.idmarque, "LibelleMarque", "")
    type_vehicule = getattr(detail.idtypevehicule, "code_type", "")
    usage = getattr(detail.idusage, "code_usage_asaci", "") or getattr(
        detail.idusage, "CodeUsage", ""
    )
    numero_police = (
        contrat.numeropolice if contrat else devis.numero_police_compagnie
    )

    return [
        "cima-jaune",
        detail.matricule,
        detail.numchassis,
        genre,
        marque,
        detail.modelevehicule,
        categorie,
        type_vehicule,
        energie,
        detail.nombreplace,
        usage,
        type_souscripteur,
        _nom_client(souscripteur),
        souscripteur.Email,
        souscripteur.Mobile or souscripteur.Telephone,
        souscripteur.CodePostal,
        _nom_client(assure),
        assure.Email,
        assure.Mobile or assure.Telephone,
        assure.CodePostal,
        _prime_rc(detail),
        numero_police,
        _date_texte(devis.dateeffet),
        _date_texte(devis.dateexpiration),
    ]


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def telecharger_demande_edition_attestation_auto(request, iddevis):
    devis = get_object_or_404(
        Devis.objects.select_related("produit", "client__idtypeclient", "assure"),
        pk=iddevis,
    )
    if devis.produit_id != 1:
        return Response(
            {"detail": "Le modèle de demande d’attestation est réservé aux devis automobiles."},
            status=status.HTTP_400_BAD_REQUEST,
        )

    details = list(
        DevisDetail.objects.filter(iddevis=devis)
        .select_related(
            "idusage",
            "idtypevehicule",
            "idgenrevehicule",
            "idmarque",
            "essence",
        )
        .prefetch_related("garanties__IdGarantie")
        .order_by("iddevisdetail")
    )
    if not details:
        return Response(
            {"detail": "Aucun véhicule n’est associé à ce devis automobile."},
            status=status.HTTP_400_BAD_REQUEST,
        )

    souscripteur = devis.client
    assure = devis.assure or souscripteur
    contrat = (
        Contrat.objects.filter(iddevis_id=devis.pk)
        .order_by("-idcontrat")
        .first()
    )
    tarifs = {
        tarif.pk: tarif
        for tarif in Tarif.objects.filter(
            pk__in={detail.idtarif for detail in details}
        ).select_related("IdCategorie")
    }
    workbook = load_workbook(MODELE_DEMANDE_ATTESTATION)
    feuille = workbook["MODELE DE DEMANDE EDITION"]

    for numero_ligne, detail in enumerate(details, start=2):
        valeurs = _ligne_attestation(
            devis, detail, contrat, souscripteur, assure, tarifs
        )
        for numero_colonne, valeur in enumerate(valeurs, start=1):
            cellule = feuille.cell(numero_ligne, numero_colonne)
            if numero_ligne > 2 and cellule.style_id == 0:
                cellule._style = copy(feuille.cell(2, numero_colonne)._style)
            cellule.value = "" if valeur is None else str(valeur)
            cellule.data_type = "s"

    buffer = io.BytesIO()
    workbook.save(buffer)
    numero = re.sub(r"[^A-Za-z0-9._-]+", "_", devis.numerodevis or str(devis.pk))
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = (
        f'attachment; filename="Demande_Edition_Attestation_{numero}.xlsx"'
    )
    return response
