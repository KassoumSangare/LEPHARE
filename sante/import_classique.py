"""
Import Santé classique (offre commerciale SANTE CLASSIQUE) : fichier d'OREOLE « FICHIER IMPORT SANTE CLASSIQUE »,
une ligne par personne assurée :

    NUMERO D'ORDRE | NUMERO DE CARTE | COLLEGE | NOM | PRENOM | DATE DE NAISSANCE | SEXE | NUMERO CMU | FILIATION |
    DATE D'ENTREE | AVENANT | PRIME NETTE ANNUELLE | PRIME NETTE PERIODIQUE | SURPRIME | PRIME NETTE TOTALE

(L'import MINENE garde le modèle d'URANUS et le barème : views_devis_sante.importer_affilies.)

Ce que le fichier renseigne à la place de l'utilisateur :
- les couvertures : COLLEGE (« 100% CI », « 80 % MONDE ENTIER », « CADRE 90% CI »…) donne la formule de l'offre
  (taux et zone) et, s'il le cite, le collège ; une couverture par valeur, créée au besoin ;
- les familles : FILIATION A (assuré principal) ouvre une famille, C (conjoint) et E (enfant) la suivent ;
- la date d'effet : la plus ancienne DATE D'ENTREE ; l'échéance un an moins un jour après ;
- les primes de l'écran (la Santé classique n'a pas de barème) : prime globale = somme des primes nettes des
  personnes hors surprime ; surprime = somme de la colonne SURPRIME (montant, ou taux si « % » : sur la prime
  annuelle) ; la majoration est l'écart entre PRIME NETTE TOTALE et prime périodique + surprime, comprise dans la
  prime globale ;
- le type de contrat : contrat société au-delà des seuils du contrat groupe (sante.groupe_adherents / affiliés).
"""

import re
import unicodedata
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from configuration_api.parametrage import valeur_parametre

ID_TARIF_CLASSIQUE = 85
ID_ZONE_COTE_IVOIRE, ID_ZONE_MONDE = 1, 2
ID_COLLEGE_PAR_DEFAUT = 10  # Famille
ID_TYPE_PARTICULIER, ID_TYPE_SOCIETE = 1, 2
LIGNES_MAX = 5000

# En-tête du fichier (normalisé : majuscules sans accents) → clé ; aide de chaque colonne (modèle et feuille Aide)
COLONNES = [
    ("NUMERO D'ORDRE", "ordre", "Facultatif : numéro de la ligne ou de la famille."),
    ("NUMERO DE CARTE", "carte", "Facultatif : numéro de carte de l'assuré (matricule)."),
    ("COLLEGE", "college", "Obligatoire : taux et zone de la couverture, ex. 100% CI, 80% MONDE ENTIER ; un collège cité "
                           "(CADRE 90% CI) est repris."),
    ("NOM", "nom", "Obligatoire."),
    ("PRENOM", "prenom", ""),
    ("DATE DE NAISSANCE", "naissance", "Obligatoire, jj/mm/aaaa."),
    ("SEXE", "sexe", "Obligatoire : M ou F."),
    ("NUMERO CMU", "cmu", ""),
    ("FILIATION", "filiation", "A assuré principal (ouvre la famille), C conjoint, E enfant ; vide : la personne est son "
                               "propre assuré principal (une famille par personne)."),
    ("DATE D'ENTREE", "entree", "Date d'entrée dans le contrat, jj/mm/aaaa ; la plus ancienne devient la date d'effet."),
    ("AVENANT", "avenant", "Facultatif : date d'entrée par avenant (incorporation), jj/mm/aaaa, ou référence de l'avenant "
                           "de la compagnie (ex. NS001/2026-00)."),
    ("PRIME NETTE ANNUELLE", "annuelle", "Prime nette de la personne pour un an."),
    ("PRIME NETTE PERIODIQUE", "periodique", "Prime nette de la période couverte (prorata) ; la prime annuelle à défaut."),
    ("SURPRIME", "surprime", "Montant, ou taux (ex. 30 %) appliqué à la prime annuelle."),
    ("PRIME NETTE TOTALE", "totale", "Prime nette à payer, surprime comprise ; prime périodique + surprime à défaut."),
]
ALIAS = {"PRENOMS": "prenom", "DATE NAISSANCE": "naissance", "LIEN": "filiation", "LIEN DE PARENTE": "filiation",
         "DATE ENTREE": "entree", "N CARTE": "carte", "NUMERO CARTE": "carte", "CMU": "cmu"}
OBLIGATOIRES = ("college", "nom", "naissance", "sexe")
# Colonnes exigées dans l'en-tête (la FILIATION peut être vide : bordereau de personnes seules, OIKOCREDIT 100 % ME)
ENTETES_REQUISES = OBLIGATOIRES + ("filiation",)
FILIATIONS = {"A": "A", "AP": "A", "ASSURE": "A", "ASSURE PRINCIPAL": "A", "ADHERENT": "A", "PRINCIPAL": "A",
              "C": "C", "CJ": "C", "CONJOINT": "C", "CONJOINTE": "C", "EPOUX": "C", "EPOUSE": "C",
              "E": "E", "EN": "E", "ENF": "E", "ENFANT": "E"}
SEXES = {"M": "M", "MASCULIN": "M", "H": "M", "HOMME": "M", "F": "F", "FEMININ": "F", "FEMME": "F"}


class ErreurImport(Exception):
    pass


def normaliser(texte):
    texte = unicodedata.normalize("NFD", str(texte or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", texte.replace("’", "'")).strip().upper()


def _date(valeur):
    if valeur in (None, ""):
        return None
    if isinstance(valeur, datetime):
        return valeur.date()
    if isinstance(valeur, date):
        return valeur
    if isinstance(valeur, (int, float)) and 1 < valeur < 80000:  # date Excel en nombre
        return date(1899, 12, 30) + timedelta(days=int(valeur))
    texte = str(valeur).strip()[:10]
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d.%m.%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(texte, fmt).date()
        except ValueError:
            continue
    raise ErreurImport(f"date illisible « {valeur} » (jj/mm/aaaa attendu)")


def _montant(valeur, pourcentage_possible=False):
    """(montant, est_un_taux) ; « 30 % » ou une cellule au format pourcentage (0,3) donnent un taux."""
    if valeur in (None, ""):
        return Decimal(0), False
    if isinstance(valeur, (int, float, Decimal)):
        return Decimal(str(valeur)), False
    texte = str(valeur).replace("\xa0", "").replace(" ", "").replace(" ", "").replace("F", "").replace(",", ".")
    taux = pourcentage_possible and texte.endswith("%")
    try:
        return Decimal(texte.rstrip("%") or "0"), taux
    except InvalidOperation as erreur:
        raise ErreurImport(f"montant illisible « {valeur} »") from erreur


def _exemple(ligne):
    """Ligne d'exemple du modèle (XXXX) : ignorée."""
    return bool(re.fullmatch(r"X+", normaliser(ligne.get("nom")).replace(" ", "")))


def _total(ligne):
    """Ligne de total du bordereau (« TOTAL » en tête, ou montants sans nom ni date de naissance) : ignorée."""
    if any(normaliser(ligne.get(cle)).startswith(("TOTAL", "SOUS-TOTAL", "SOUS TOTAL")) for cle in ("ordre", "carte", "college", "nom")):
        return True
    return ligne.get("nom") in (None, "") and ligne.get("naissance") in (None, "")


def _date_ou_reference(valeur):
    """AVENANT : date d'entrée par avenant, ou référence de l'avenant de la compagnie (NS001/2026-00) → (date, référence)."""
    try:
        return _date(valeur), ""
    except ErreurImport:
        return None, str(valeur).strip()


def lire(fichier):
    """Lignes du fichier : [{_ligne, clés de COLONNES (valeurs brutes), _surprime_taux}], en-tête en 1re ligne."""
    classeur = load_workbook(fichier, data_only=True, read_only=False)
    feuille = classeur.worksheets[0]
    connues = {normaliser(titre): cle for titre, cle, _ in COLONNES} | ALIAS
    entetes = {}
    for cellule in feuille[1]:
        cle = connues.get(normaliser(cellule.value))
        if cle and cle not in entetes:
            entetes[cle] = cellule.column
    manquantes = [titre for titre, cle, _ in COLONNES if cle in ENTETES_REQUISES and cle not in entetes]
    if manquantes:
        raise ErreurImport("Ce fichier n'est pas le modèle d'import Santé classique : colonne"
                           f"{'s' if len(manquantes) > 1 else ''} {', '.join(manquantes)} absente{'s' if len(manquantes) > 1 else ''} "
                           "de la 1re ligne.")
    lignes = []
    for numero in range(2, feuille.max_row + 1):
        ligne = {"_ligne": numero}
        for cle, colonne in entetes.items():
            cellule = feuille.cell(numero, colonne)
            ligne[cle] = cellule.value
            if cle == "surprime":
                ligne["_surprime_taux"] = "%" in str(cellule.number_format or "") and cellule.value not in (None, "")
        if all(v in (None, "") for k, v in ligne.items() if not k.startswith("_")) or _exemple(ligne) or _total(ligne):
            continue
        lignes.append(ligne)
        if len(lignes) > LIGNES_MAX:
            raise ErreurImport(f"{LIGNES_MAX} personnes au plus par fichier.")
    if not lignes:
        raise ErreurImport("Le fichier ne contient aucune personne (lignes 2 et suivantes).")
    return lignes


def formules_classiques(cursor):
    """{(taux, zone): (idoffre, libellé)} des formules de l'offre Santé classique (« … 100% COTE D'IVOIRE »)."""
    cursor.execute("SELECT idoffre, libelleoffre FROM stdoffre WHERE idtarif = %s AND actif", [ID_TARIF_CLASSIQUE])
    resultat = {}
    for id_offre, libelle in cursor.fetchall():
        texte = normaliser(libelle)
        taux = re.search(r"(\d+(?:[.,]\d+)?)\s*%", texte)
        if taux:
            zone = ID_ZONE_MONDE if "MONDE" in texte else ID_ZONE_COTE_IVOIRE
            resultat[(Decimal(taux.group(1).replace(",", ".")), zone)] = (id_offre, libelle)
    return resultat


def colleges(cursor):
    cursor.execute("SELECT idcollege, libellecollege FROM stdcollege WHERE idcollege BETWEEN 1 AND 10")
    # Libellés les plus longs d'abord : « CADRE SUPERIEUR » avant « CADRE »
    return sorted(((i, normaliser(l)) for i, l in cursor.fetchall()), key=lambda x: -len(x[1]))


def couverture_du_college(texte, formules, liste_colleges, id_college_defaut):
    """COLLEGE du fichier → (idoffre, libellé de la formule, idzone, idcollege)."""
    valeur = normaliser(texte)
    taux = re.search(r"(\d+(?:[.,]\d+)?)\s*%", valeur)
    if not taux:
        raise ErreurImport(f"collège « {texte} » : taux de couverture absent (ex. 100% CI)")
    zone = ID_ZONE_MONDE if re.search(r"MONDE|\bMOE\b|\bME\b|INTERNATIONAL", valeur) else ID_ZONE_COTE_IVOIRE
    formule = formules.get((Decimal(taux.group(1).replace(",", ".")), zone))
    if not formule:
        raise ErreurImport(f"collège « {texte} » : aucune formule Santé classique à {taux.group(1)} % "
                           f"{'Monde entier' if zone == ID_ZONE_MONDE else 'Côte d’Ivoire'}")
    college = next((i for i, l in liste_colleges if re.search(rf"\b{re.escape(l)}\b", valeur)), id_college_defaut)
    return formule[0], formule[1], zone, college


def analyser(cursor, fichier, id_college_defaut=ID_COLLEGE_PAR_DEFAUT):
    """Lecture et contrôle du fichier, sans rien écrire : familles par couverture, personnes prêtes à importer
    (lignes au format de sante.utils.import_insured), primes et dates déduites, erreurs par ligne."""
    lignes = lire(fichier)
    formules, liste_colleges = formules_classiques(cursor), colleges(cursor)
    erreurs, avertissements, familles, personnes = [], [], [], 0
    totaux = {"annuelle": Decimal(0), "base": Decimal(0), "surprime": Decimal(0), "majoration": Decimal(0), "reference": Decimal(0)}
    entrees, sans_filiation = [], 0
    for ligne in lignes:
        try:
            for cle in OBLIGATOIRES:
                if ligne.get(cle) in (None, ""):
                    raise ErreurImport(f"{next(t for t, c, _ in COLONNES if c == cle)} obligatoire")
            # FILIATION vide : la personne est son propre assuré principal (une famille par personne)
            if ligne.get("filiation") in (None, ""):
                filiation = "A"
                sans_filiation += 1
            else:
                filiation = FILIATIONS.get(normaliser(ligne["filiation"]))
            if not filiation:
                raise ErreurImport(f"filiation « {ligne['filiation']} » : A, C ou E")
            sexe = SEXES.get(normaliser(ligne["sexe"]))
            if not sexe:
                raise ErreurImport(f"sexe « {ligne['sexe']} » : M ou F")
            naissance = _date(ligne["naissance"])
            if not naissance:
                raise ErreurImport("date de naissance obligatoire")
            entree = _date(ligne.get("entree"))
            avenant, _ = _date_ou_reference(ligne.get("avenant"))
            annuelle, _ = _montant(ligne.get("annuelle"))
            periodique, _ = _montant(ligne.get("periodique"))
            surprime, taux = _montant(ligne.get("surprime"), pourcentage_possible=True)
            totale, _ = _montant(ligne.get("totale"))
            reference = periodique or annuelle
            if taux or ligne.get("_surprime_taux"):
                pourcentage = surprime if taux else surprime * 100
                surprime = ((annuelle or periodique) * pourcentage / 100).quantize(Decimal(1))
            base = (totale - surprime) if totale else reference
            if base < 0:
                raise ErreurImport("prime nette totale inférieure à la surprime")
            couverture = couverture_du_college(ligne["college"], formules, liste_colleges, id_college_defaut)
        except ErreurImport as erreur:
            erreurs.append({"Ligne": ligne["_ligne"], "Nom": f"{ligne.get('nom') or ''} {ligne.get('prenom') or ''}".strip(),
                            "Message": str(erreur)})
            continue
        if filiation == "A":
            familles.append({"couverture": couverture, "lignes": []})
        elif not familles:
            erreurs.append({"Ligne": ligne["_ligne"], "Nom": f"{ligne['nom']} {ligne.get('prenom') or ''}".strip(),
                            "Message": "affilié avant tout assuré principal (filiation A)"})
            continue
        famille = familles[-1]
        if couverture[:3] != famille["couverture"][:3]:
            avertissements.append({"Ligne": ligne["_ligne"], "Nom": f"{ligne['nom']} {ligne.get('prenom') or ''}".strip(),
                            "Message": f"collège « {ligne['college']} » différent de celui de son assuré principal : "
                                       "famille rattachée à la couverture de l'assuré principal"})
        personnes += 1
        if entree or avenant:
            entrees.append(entree or avenant)
        totaux["annuelle"] += annuelle
        totaux["base"] += base
        totaux["surprime"] += surprime
        totaux["reference"] += reference
        totaux["majoration"] += max(base - reference, Decimal(0)) if totale else Decimal(0)
        iso = lambda d: d.isoformat() if d else "NA"  # noqa: E731
        famille["lignes"].append({
            "_ligne": ligne["_ligne"], "nom": str(ligne["nom"]).strip(), "prenoms": str(ligne.get("prenom") or "").strip() or "NA",
            "numerocni": "NA", "datenaissance": naissance.isoformat(), "numerocmu": str(ligne.get("cmu") or "").strip() or "NA",
            "sexe": sexe, "numerotelephone": "NA", "lienparente": filiation, "dateeffet": iso(entree or avenant),
            "dateentree": iso(avenant or entree), "adresseadherent": "NA", "emailadherent": "NA", "nombreaffection": "0",
            "matricule": str(ligne.get("carte") or "").strip() or "NA", "groupesanguin": "",
            "montantsurprime": str(surprime), "_base": base,
        })
    if sans_filiation:
        avertissements.insert(0, {"Ligne": "", "Nom": "", "Message": (
            f"FILIATION vide sur {sans_filiation} ligne{'s' if sans_filiation > 1 else ''} : chaque personne concernée est "
            "prise comme assuré principal de sa propre famille.")})
    # Couvertures : une par formule, zone et collège
    couvertures = {}
    for famille in familles:
        cle = famille["couverture"]
        c = couvertures.setdefault(cle, {"IdOffre": cle[0], "Formule": cle[1], "IdZone": cle[2], "IdCollege": cle[3],
                                         "Familles": 0, "Personnes": 0, "PrimeNette": Decimal(0), "Surprime": Decimal(0)})
        c["Familles"] += 1
        c["Personnes"] += len(famille["lignes"])
        c["PrimeNette"] += sum((l["_base"] for l in famille["lignes"]), Decimal(0))
        c["Surprime"] += sum((Decimal(l["montantsurprime"]) for l in famille["lignes"]), Decimal(0))
    noms_colleges = dict((i, l) for i, l in liste_colleges)
    effet = min(entrees) if entrees else None
    seuil_adherents = int(valeur_parametre("sante.groupe_adherents", 5))
    seuil_personnes = int(valeur_parametre("sante.groupe_affilies", 25))
    groupe = len(familles) > seuil_adherents or personnes > seuil_personnes
    reference = totaux["reference"]
    return {
        "Familles": len(familles), "Personnes": personnes, "Erreurs": erreurs, "Avertissements": avertissements,
        "Couvertures": [{**c, "College": noms_colleges.get(c["IdCollege"], "").title(), "PrimeNette": float(c["PrimeNette"]),
                         "Surprime": float(c["Surprime"])} for c in couvertures.values()],
        "DateEffet": effet.isoformat() if effet else None,
        "DateExpiration": (date(effet.year + 1, effet.month, effet.day if not (effet.month == 2 and effet.day == 29) else 28)
                           - timedelta(days=1)).isoformat() if effet else None,
        "PrimeGlobale": float(totaux["base"].quantize(Decimal(1))), "MontantSurprime": float(totaux["surprime"].quantize(Decimal(1))),
        "PrimeAnnuelle": float(totaux["annuelle"]), "Majoration": float(totaux["majoration"].quantize(Decimal(1))),
        "TauxMajoration": round(float(totaux["majoration"] * 100 / reference), 2) if reference else 0.0,
        "TypeContrat": ID_TYPE_SOCIETE if groupe else ID_TYPE_PARTICULIER, "Groupe": groupe,
        "_familles": familles,
    }


def modele():
    """Classeur modèle : en-tête d'OREOLE, une ligne d'exemple, feuille d'aide."""
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Feuil1"
    feuille.append([titre for titre, _, _ in COLONNES])
    feuille.append([1, "", "100% CI", "XXXXXXXXXX", "XXXXXXXX", "XX/XX/XXXX", "M", "", "A", date(2024, 2, 5), None, None, None, None, None])
    for cellule in feuille[1]:
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = PatternFill("solid", fgColor="1F4E79")
    for colonne, largeur in zip("ABCDEFGHIJKLMNO", (12, 16, 14, 22, 22, 16, 8, 16, 10, 14, 12, 18, 18, 12, 18)):
        feuille.column_dimensions[colonne].width = largeur
    for cellule in ("F2", "J2", "K2"):
        feuille[cellule].number_format = "dd/mm/yyyy"
    aide = classeur.create_sheet("Aide")
    aide.append(["Colonne", "Contenu"])
    for titre, _, texte in COLONNES:
        aide.append([titre, texte])
    aide.append([])
    aide.append(["Familles", "Chaque assuré principal (A) est suivi de son conjoint (C) et de ses enfants (E). La ligne d'exemple "
                             "(XXXX) est ignorée."])
    aide.column_dimensions["A"].width = 26
    aide.column_dimensions["B"].width = 110
    for cellule in aide[1]:
        cellule.font = Font(bold=True)
    return classeur
